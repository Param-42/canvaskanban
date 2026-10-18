#!/usr/bin/env python3
"""
Builds the CanvasKanban project report for the college assignment.

Two files are produced next to this script:

    CanvasKanban-Project-Report.pdf    the document to hand in
    CanvasKanban-Project-Report.docx   the same document, editable in Word /
                                       Google Docs

Everything that has to be filled in by the team is written between square
brackets, for example [Member 1 - Full Name]. The .docx marks all of those in
yellow and the .pdf prints them in dark red, so nothing gets forgotten.

The numbers in section "Project statistics" were taken from the repository
itself (file and line counts, test count, release assets).

Usage (needs python-docx and reportlab):

    python3 -m venv .venv
    .venv/bin/pip install python-docx reportlab
    .venv/bin/python build_report.py
"""

import os
import re

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak,
                                Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
SHOTS = os.path.join(REPO, "docs", "screenshots")

PLACEHOLDER = re.compile(r"(\[[^\]]+\])")

# ---------------------------------------------------------------------------
# 1. The content. Edit the text here, then run the script again.
# ---------------------------------------------------------------------------

TITLE = "CanvasKanban - A Desktop Kanban Board Application"
SUBTITLE = "Software Project Report"

SUBMITTED = [
    ("Course", "[Course name, e.g. Object Oriented Programming with Java]"),
    ("College / University", "[College or university name]"),
    ("Department", "[Department]"),
    ("Guide / Faculty", "[Name of the faculty member]"),
    ("Academic year", "[2026 - 2027]"),
    ("Submission date", "[Date of submission]"),
]

DESCRIPTION = (
    "CanvasKanban is a single-user desktop Kanban board written in Java 21 and JavaFX. Boards, "
    "tasks and settings are kept in an embedded SQLite database inside the user's own data "
    "folder; there is no account, no server and no network access at all. A board consists of "
    "the three columns <b>To Do</b>, <b>In Progress</b> and <b>Done</b>. Tasks carry a title, a "
    "description, a priority (Low / Medium / High), the date they were created and an optional "
    "due date, and they can be dragged into another column to change their status or dragged up "
    "and down to reorder them. Six colour themes can be chosen per board and are remembered. The "
    "application is delivered as self-contained packages - MSI and ZIP for Windows, DEB, RPM and "
    "tar.gz for Linux - that bring their own Java runtime with them, so the user does not have to "
    "install Java. Two one-command installers (install.sh and install.ps1) download and install "
    "the newest release, create the menu or Start-menu entries and can remove everything again."
)

OBJECTIVES = [
    "Design and build a complete, working desktop application with a clean layered "
    "architecture (user interface &gt; service &gt; data access objects &gt; database).",
    "Practise persistent storage with a relational database: an SQLite schema, prepared "
    "statements and transactions, without hiding the SQL behind an ORM.",
    "Implement an interactive JavaFX interface including drag and drop between columns, "
    "dialog-based create/edit flows and a CSS based theming system.",
    "Apply professional engineering practices: a build with Maven, an automated test suite, "
    "continuous integration on two operating systems and versioned releases.",
    "Ship an application that a non-technical user can install in a single step, without "
    "installing Java or any other dependency first.",
]

TECHNOLOGIES = [
    ("Layer / Area", "Technology", "What it is used for"),
    ("Language &amp; runtime", "Java 21 (Temurin JDK)",
     "All application code; a trimmed private runtime is produced with jlink"),
    ("Frontend (GUI)", "JavaFX 21 (controls, graphics, base)",
     "The whole interface - windows, board, cards, dialogs and drag and drop; built by hand, "
     "without FXML"),
    ("Frontend styling", "CSS (JavaFX CSS, custom -app-* variables)",
     "All colours, borders and spacing; six theme files on top of one base file"),
    ("Backend (logic)", "Plain Java service layer + DAO pattern",
     "Business rules (BoardService) and all SQL (BoardDao, TaskDao, SettingDao)"),
    ("Database", "SQLite, via the xerial sqlite-jdbc driver",
     "Embedded file database with the tables boards, tasks and settings"),
    ("Build", "Apache Maven 3.9 (compiler, jar, dependency, surefire, javafx-maven-plugin)",
     "Compiling, dependency management, running the app and producing the JAR"),
    ("Testing", "JUnit 5 (67 tests)",
     "Database, DAO, service, theme and real user-interface tests"),
    ("Packaging", "jpackage and jlink (part of the JDK)",
     "Application image plus native installers: MSI, DEB, RPM, tar.gz, ZIP"),
    ("Continuous integration", "GitHub Actions",
     "Ubuntu and Windows matrix builds, a Fedora 43 container job for the RPM, a Windows job "
     "with the WiX Toolset for the MSI, and the release job"),
    ("Distribution", "Bash (install.sh) and PowerShell (install.ps1)",
     "One-command install, update and uninstall on Linux and Windows"),
    ("Logging", "SLF4J (slf4j-nop)",
     "A single logging facade, quiet by default"),
    ("Version control", "Git and GitHub",
     "Source hosting, issues, pull-request style review and releases"),
]

FEATURES = [
    "Boards can be created, renamed and deleted; a fresh installation starts with one example "
    "board.",
    "Tasks with title, description, priority (Low / Medium / High), creation date and an "
    "optional due date.",
    "The three fixed columns To Do, In Progress and Done.",
    "Drag and drop: a card dragged into another column changes its status; dragging it up or "
    "down inside a column reorders it, with a blue line showing where it will land.",
    "Right-click menu on a card to delete the task (with confirmation) or to read the full "
    "description.",
    "Six built-in themes (Minimal Light, Dark Mode, Midnight Blue, Forest Green, Lavender, Warm "
    "Sand); the choice is stored per board.",
    "Everything is stored locally in SQLite; the last opened board and its theme are restored on "
    "the next start.",
    "A status bar that reports the last action, the location of the database and the number of "
    "tasks on the board.",
    "Self-contained installers for Windows and Linux plus one-command installers that need "
    "nothing but curl or PowerShell.",
    "A suite of 67 automated tests and a continuous integration pipeline that builds and tests "
    "the application on Ubuntu and Windows on every push.",
]

REPOSITORY = "https://github.com/Param-42/canvaskanban"
REPOSITORY_NOTE = (
    "The repository contains the full source code, the build and packaging scripts, the "
    "continuous integration workflow, the screenshots and the user and developer "
    "documentation."
)

TEAM = [
    ("1", "[Member 1 - Full Name]", "[Enrollment Number]", "Team Leader"),
    ("2", "[Member 2 - Full Name]", "[Enrollment Number]", "Team Member"),
    ("3", "[Member 3 - Full Name]", "[Enrollment Number]", "Team Member"),
    ("4", "[Member 4 - Full Name]", "[Enrollment Number]", "Team Member"),
    ("5", "[Member 5 - Full Name]", "[Enrollment Number]", "Team Member"),
]

CONTRIBUTIONS = [
    {
        "name": "[Member 1 - Full Name]",
        "enroll": "[Enrollment Number]",
        "role": "Team Leader - planning, architecture, build, continuous integration and delivery",
        "tasks": [
            "Set up the GitHub repository, the Maven project and the package structure, and "
            "defined the layered architecture (ui &gt; service &gt; dao &gt; database) together with "
            "the coding conventions the team followed.",
            "Split the work into modules, kept track of progress, reviewed every change before it "
            "reached the main branch and resolved the conflicts that came up.",
            "Integrated the five modules into one application and led the debugging of the "
            "defects that only appeared once they were combined (board switching, task ordering, "
            "database locking).",
            "Configured the build in <b>pom.xml</b>: Java 21, the JavaFX and SQLite dependencies, "
            "the JUnit 5 test setup and the jar, dependency, surefire and javafx-maven plugins.",
            "Wrote the packaging scripts <b>package.sh</b> and <b>package.bat</b>, which drive "
            "jpackage to produce the application image and the native installers (MSI, DEB, RPM, "
            "tar.gz).",
            "Built the continuous integration pipeline in <b>.github/workflows/build.yml</b>: an "
            "Ubuntu and Windows matrix that compiles and tests, a Fedora 43 container job for the "
            "RPM, a Windows job that installs the WiX Toolset for the MSI, and a release job that "
            "attaches the packages to a GitHub release with stable file names.",
            "Managed the releases, wrote the one-command installers <b>install.sh</b> and "
            "<b>install.ps1</b>, and verified them against the real releases.",
        ],
        "modules": [
            "pom.xml - Maven build and dependency configuration",
            "package.sh, package.bat - jpackage packaging for Linux and Windows",
            ".github/workflows/build.yml - the whole CI/CD pipeline",
            "install.sh, install.ps1 - one-command install, update and uninstall",
        ],
        "tech": "Java 21, Maven, jpackage, jlink, GitHub Actions, YAML, Bash, PowerShell, WiX "
                "Toolset, Fedora container builds, Git and GitHub releases.",
    },
    {
        "name": "[Member 2 - Full Name]",
        "enroll": "[Enrollment Number]",
        "role": "Team Member - database and persistence",
        "tasks": [
            "Designed the SQLite schema (the tables <b>boards</b>, <b>tasks</b> and <b>settings</b>, "
            "with the appropriate indexes and foreign keys) and the behaviour that creates missing "
            "tables when the application starts.",
            "Implemented the connection layer: one JDBC connection per call, opened and closed "
            "inside the code, with prepared statements everywhere so no SQL is ever built by "
            "string concatenation.",
            "Implemented the platform-independent data folder lookup - the XDG data directory on "
            "Linux, %APPDATA% on Windows and Application Support on macOS - including the "
            "CANVASKANBAN_DATA_DIR environment variable and the -Dcanvaskanban.data.dir system "
            "property that move the database elsewhere.",
            "Implemented the date helpers (ISO text to LocalDate and back) and the "
            "DataAccessException wrapper that keeps JDBC details out of the user interface.",
            "Wrote the database test class and supported the DAO tests with test databases.",
        ],
        "modules": [
            "src/main/java/com/canvaskanban/database/AppPaths.java - where the database lives on "
            "each operating system",
            ".../database/Database.java - connections, queries, transactions",
            ".../database/Schema.java - tables, indexes and migration on start",
            ".../database/SqlDates.java, .../database/DataAccessException.java",
            "src/test/java/com/canvaskanban/database/DatabaseTest.java",
        ],
        "tech": "Java 21, JDBC, SQLite (sqlite-jdbc), SQL DDL, prepared statements, transactions, "
                "JUnit 5.",
    },
    {
        "name": "[Member 3 - Full Name]",
        "enroll": "[Enrollment Number]",
        "role": "Team Member - data access layer and business rules",
        "tasks": [
            "Implemented <b>BoardDao</b>, <b>TaskDao</b> and <b>SettingDao</b> in plain SQL, "
            "without an ORM: create, read, update and delete for boards and tasks, ordered "
            "queries and a key/value store for settings such as the last open board and the theme "
            "of each board.",
            "Implemented moving and reordering tasks in a single transaction, rewriting the "
            "positions of both the source and the target column so that the order in the database "
            "can never become inconsistent.",
            "Implemented <b>BoardService</b>, the only interface the user interface is allowed to "
            "use. It carries the rules of the application: every board gets the three columns, "
            "deleting a board deletes its tasks, input is validated and trimmed, and positions "
            "are assigned when a task is created.",
            "Implemented the domain model (Board, Task, Priority, TaskStatus, Theme) that is "
            "passed between the layers.",
            "Wrote the largest part of the automated tests: the DAO tests and the service tests, "
            "including the cases for moving a task between columns and reordering inside a "
            "column.",
        ],
        "modules": [
            "src/main/java/com/canvaskanban/dao/BoardDao.java, TaskDao.java, SettingDao.java",
            ".../service/BoardService.java - the business rules of the application",
            ".../model/Board.java, Task.java, Priority.java, TaskStatus.java, Theme.java",
            "src/test/java/com/canvaskanban/dao/* and .../service/BoardServiceTest.java",
        ],
        "tech": "Java 21, JDBC, SQL, transactions, layered design, JUnit 5, SLF4J.",
    },
    {
        "name": "[Member 4 - Full Name]",
        "enroll": "[Enrollment Number]",
        "role": "Team Member - JavaFX user interface",
        "tasks": [
            "Implemented the two entry points: <b>Launcher</b> (used by the packaged application, "
            "the JAR and the tests, because Java refuses to start an Application subclass from "
            "the class path) and <b>MainApp</b> (the JavaFX lifecycle: init, start, stop).",
            "Built the main screen in <b>MainView</b>: the toolbar with the board selector and the "
            "actions, the board area with the three columns, and the status bar in the bottom "
            "left and right corners.",
            "Implemented the dialogs in <b>Dialogs</b>: creating and editing a task (title, "
            "description, priority, due date), creating and renaming a board, and the "
            "confirmation alerts for deleting a task or a board.",
            "Implemented the task card in <b>TaskCardView</b>: title, description preview, the "
            "coloured priority badge, the due date and the right-click menu (delete task, show "
            "description).",
            "Connected the interface to BoardService, so that the board is rebuilt whenever "
            "something changed and the selected board stays in sync with the database.",
            "Implemented the user feedback: the action messages ('Task deleted', 'Board created') "
            "and the footer that shows the database location and the number of tasks.",
            "Contributed to the user-interface test class, which builds the real scene graph in "
            "the test and checks the columns, the cards, the dialogs and the move operations.",
        ],
        "modules": [
            "src/main/java/com/canvaskanban/Launcher.java, MainApp.java",
            ".../ui/MainView.java - toolbar, board area, status bar",
            ".../ui/Dialogs.java - the create/edit/delete dialogs and alerts",
            ".../ui/TaskCardView.java - one task card",
            "src/test/java/com/canvaskanban/ui/MainViewTest.java",
        ],
        "tech": "Java 21, JavaFX 21 (controls, scene graph, layouts, dialogs, no FXML), CSS, JUnit "
                "5 with a real JavaFX environment (Xvfb in continuous integration).",
    },
    {
        "name": "[Member 5 - Full Name]",
        "enroll": "[Enrollment Number]",
        "role": "Team Member - drag and drop, themes and documentation",
        "tasks": [
            "Implemented the drag-and-drop behaviour: <b>BoardColumnView</b> accepts cards from "
            "other columns (status change) and from its own column (reordering), shows a "
            "highlight while a card hovers over it and draws the blue indicator line at the exact "
            "position where the card will land.",
            "Implemented the theming system: <b>base.css</b> contains all of the layout and reads "
            "its colours from -app-* variables, and each of the six themes is a small file that "
            "only sets those variables, so a theme can never break the layout.",
            "Implemented <b>ThemeManager</b> (which stylesheet belongs to a theme and how it is "
            "applied) and <b>ThemeListCell</b>, which draws the coloured dots in the theme "
            "selector.",
            "Implemented per-board theme persistence, so that each board opens again with the look "
            "it had.",
            "Wrote the theme tests, which check that every theme file exists and defines the "
            "variables the base file needs, and produced the screenshots used in the README.",
            "Wrote and maintained the documentation: the README and docs/USING.md, "
            "docs/BUILDING.md and docs/ARCHITECTURE.md, including the installation instructions "
            "for end users and for developers.",
        ],
        "modules": [
            "src/main/java/com/canvaskanban/ui/BoardColumnView.java - drop targets and the "
            "indicator line",
            ".../ui/ThemeManager.java, .../ui/ThemeListCell.java",
            "src/main/resources/css/base.css plus the six theme files (minimal-light, dark-mode, "
            "midnight-blue, forest-green, lavender, warm-sand)",
            "src/test/java/com/canvaskanban/ui/ThemeCssTest.java",
            "README.md and the docs/ folder (USING.md, BUILDING.md, ARCHITECTURE.md)",
        ],
        "tech": "Java 21, the JavaFX drag-and-drop API, CSS with custom variables (JavaFX CSS), "
                "JUnit 5, Markdown, screenshot preparation.",
    },
]

SHARED = [
    "<b>[Member 3] and [Member 4]</b> tested the board flows together: every change in the user "
    "interface was checked against the database layer, which is how the ordering bugs around "
    "drag and drop were found.",
    "<b>[Member 4] and [Member 5]</b> worked together on the interaction between drag and drop "
    "and the themes, so that the drop indicator and the column highlight stay readable in all "
    "six themes, including the dark ones.",
    "<b>[Member 2] and [Member 3]</b> agreed on the contract between the schema and the DAOs "
    "(column names, types, ordering) before either of them wrote code, then reviewed each "
    "other's SQL.",
    "<b>[Member 1] and [Member 3]</b> worked together on integration: the service layer was "
    "adjusted several times while the packaging and the tests were set up.",
    "<b>All five members</b> took part in testing the packaged builds - installing the MSI, the "
    "DEB and the RPM and running the application on their own machines - and in reviewing the "
    "documentation before submission.",
]

MATRIX_HEADER = ["Module / Area", "Member 1", "Member 2", "Member 3", "Member 4", "Member 5"]
MATRIX = [
    ("Planning, architecture, code review, integration", "Owner", "Support", "Support", "Support",
     "Support"),
    ("Maven build configuration (pom.xml)", "Owner", "-", "-", "-", "-"),
    ("SQLite schema, connections and data paths", "Support", "Owner", "Support", "-", "-"),
    ("DAO layer and business service", "-", "Support", "Owner", "-", "-"),
    ("JavaFX interface: shell, dialogs, cards", "Support", "-", "Support", "Owner", "Support"),
    ("Drag and drop between columns and reordering", "-", "-", "Support", "Support", "Owner"),
    ("Themes and CSS styling", "-", "-", "-", "Support", "Owner"),
    ("Automated tests (67 in total)", "Support", "Owner (DB)", "Owner (logic)",
     "Owner (UI)", "Owner (themes)"),
    ("Packaging with jpackage (MSI, DEB, RPM, tar.gz)", "Owner", "-", "-", "-", "Support"),
    ("Continuous integration and releases", "Owner", "-", "Support", "-", "Support"),
    ("One-command installers (install.sh, install.ps1)", "Owner", "-", "-", "-", "-"),
    ("Documentation and screenshots", "Support", "Support", "Support", "Support", "Owner"),
]

STATISTICS = [
    "Application code: <b>22 Java files with 2,627 lines</b> in src/main/java (user interface, "
    "service, DAO, database, model).",
    "Test code: <b>7 test classes with 1,255 lines and 67 test methods</b> in src/test/java; the "
    "suite covers the database layer, the DAOs, the service rules, the theme files and the real "
    "user interface.",
    "Styling: <b>7 CSS files with 664 lines</b> - one base file that holds all layout and six "
    "themes that only define colours.",
    "Database: <b>3 tables</b> (boards, tasks, settings) in a single SQLite file inside the user's "
    "data folder - no server and no configuration.",
    "Continuous integration: every push runs <b>mvn clean package</b> and the tests on "
    "<b>ubuntu-latest and windows-latest</b>; the user-interface tests really execute on the "
    "build server because a virtual screen (Xvfb) is installed.",
    "Release <b>v1.0.0</b> ships <b>6 packages</b>: Windows MSI and ZIP, Linux DEB, RPM, tar.gz "
    "and the plain JAR, all attached to the GitHub release.",
    "The installers were exercised on real machines and inside the CI containers (the RPM is "
    "built in a Fedora 43 container, the MSI on a Windows runner with the WiX Toolset).",
]

SCREENSHOTS = [
    ("minimal-light.png", "The board with the Minimal Light theme: three columns, task cards "
                          "with priority badges and due dates, toolbar and status bar."),
    ("dark-mode.png", "The same board with the Dark Mode theme - the theme is stored per board."),
    ("drag-and-drop-result.png", "The board after several drag-and-drop moves, an edit and a "
                                 "deletion."),
    ("theme-selector.png", "The theme selector with the coloured dots of the six themes."),
]

# ---------------------------------------------------------------------------
# 2. The PDF
# ---------------------------------------------------------------------------

def _pdf_markup(text):
    """Prints the [to be filled in] parts in dark red so they are easy to find."""
    return PLACEHOLDER.sub(r'<font color="#B00020"><b>\1</b></font>', text)


def build_pdf(path):
    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=24, leading=29,
                           spaceAfter=6, alignment=TA_CENTER)
    subtitle = ParagraphStyle("st", parent=styles["Normal"], fontSize=14, leading=18,
                              alignment=TA_CENTER, textColor=colors.HexColor("#444444"))
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=16, leading=20,
                        spaceBefore=14, spaceAfter=8, textColor=colors.HexColor("#12324f"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13, leading=17,
                        spaceBefore=10, spaceAfter=6, textColor=colors.HexColor("#1d4e74"))
    h3 = ParagraphStyle("h3", parent=styles["Heading3"], fontSize=11.5, leading=15,
                        spaceBefore=9, spaceAfter=4, textColor=colors.HexColor("#2d6f9f"))
    body = ParagraphStyle("b", parent=styles["BodyText"], fontSize=10.5, leading=15,
                          alignment=TA_LEFT, spaceAfter=5)
    small = ParagraphStyle("s", parent=body, fontSize=9, leading=12.5,
                           textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("c", parent=body, fontSize=9.5, leading=13, spaceAfter=0)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Helvetica-Bold")

    def P(text, style=body):
        return Paragraph(_pdf_markup(text), style)

    def bullet_list(items, style=body):
        return ListFlowable(
            [ListItem(P(i, style)) for i in items],
            bulletType="bullet", bulletFontName="Helvetica", bulletFontSize=7,
            bulletOffsetY=-1, leftIndent=14, spaceAfter=6)

    def table(rows, widths, bold_first_col=False, bold_header=True):
        data = []
        for r, row in enumerate(rows):
            line = []
            for c, value in enumerate(row):
                style = cellb if (r == 0 and bold_header) or (bold_first_col and c == 0) else cell
                line.append(Paragraph(_pdf_markup(str(value)), style))
            data.append(line)
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef4")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b9c6d2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        if bold_first_col:
            t.setStyle(TableStyle([("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold")]))
        return t

    story = []

    # -- title page ---------------------------------------------------------
    story.append(Spacer(1, 3.2 * cm))
    story.append(Paragraph(_pdf_markup(TITLE), title))
    story.append(Paragraph(_pdf_markup(SUBTITLE), subtitle))
    story.append(Spacer(1, 1.6 * cm))
    story.append(table([(k, v) for k, v in SUBMITTED], [4.6 * cm, 11.4 * cm]))
    story.append(Spacer(1, 1.2 * cm))
    story.append(P("Submitted by", h3))
    story.append(table([("S. No.", "Full name", "Enrollment number", "Role")] + TEAM,
                       [1.4 * cm, 5.6 * cm, 4.6 * cm, 4.4 * cm]))
    story.append(PageBreak())

    # -- 1. project information --------------------------------------------
    story.append(P("1. Project information", h1))

    story.append(P("1.1 Project title", h2))
    story.append(P("CanvasKanban - A Desktop Kanban Board Application", body))

    story.append(P("1.2 Project description", h2))
    story.append(P(DESCRIPTION, body))

    story.append(P("1.3 Project objectives", h2))
    story.append(bullet_list(OBJECTIVES))

    story.append(P("1.4 Technologies used", h2))
    story.append(table(TECHNOLOGIES, [3.5 * cm, 5.2 * cm, 7.3 * cm]))

    story.append(P("1.5 Key features and functionalities", h2))
    story.append(bullet_list(FEATURES))

    story.append(P("1.6 GitHub repository", h2))
    story.append(P('<font color="#0b5cad">%s</font>' % REPOSITORY, body))
    story.append(P(REPOSITORY_NOTE, body))
    story.append(PageBreak())

    # -- 2. team ------------------------------------------------------------
    story.append(P("2. Team members and responsibilities", h1))
    story.append(P("The team has five members. One member acts as the team leader and is "
                   "responsible for planning, architecture, integration and delivery; the other "
                   "four members own one module each. Every member wrote code, tests and "
                   "documentation, and everybody reviewed the work of at least one other "
                   "member.", body))
    story.append(Spacer(1, 0.3 * cm))
    story.append(table([("S. No.", "Full name", "Enrollment number", "Role")] + TEAM,
                       [1.4 * cm, 5.6 * cm, 4.6 * cm, 4.4 * cm]))

    # -- 3. contributions ---------------------------------------------------
    story.append(P("3. Individual contributions", h1))
    for i, m in enumerate(CONTRIBUTIONS, start=1):
        block = [P("3.%d %s" % (i, m["name"]), h2),
                 P("<b>Enrollment number:</b> %s" % m["enroll"], small),
                 P("<b>Role in the project:</b> %s" % m["role"], small),
                 P("<b>Specific tasks completed by the member:</b>", cell),
                 bullet_list(m["tasks"])]
        story.append(KeepTogether(block))
        story.append(P("<b>Features / modules developed by the member:</b>", cell))
        story.append(bullet_list(m["modules"]))
        story.append(P("<b>Technologies and tools used by the member:</b> %s" % m["tech"], body))
        story.append(Spacer(1, 0.25 * cm))

    # -- 4. shared work -----------------------------------------------------
    story.append(P("4. Contributions shared by two or more members", h1))
    story.append(bullet_list(SHARED))

    # -- 5. matrix ----------------------------------------------------------
    story.append(P("5. Contribution matrix", h1))
    story.append(P("'Owner' means the member designed and implemented that part; 'Support' means "
                   "the member contributed to it, for example through pair work, testing or "
                   "review.", small))
    story.append(Spacer(1, 0.2 * cm))
    story.append(table([MATRIX_HEADER] + [list(r) for r in MATRIX],
                       [5.8 * cm, 2.1 * cm, 2.1 * cm, 2.1 * cm, 2.1 * cm, 2.1 * cm]))

    # -- 6. statistics ------------------------------------------------------
    story.append(P("6. Project statistics and evidence", h1))
    story.append(P("The numbers below were taken from the repository and from the continuous "
                   "integration runs, so they can be checked by anyone with access to it.", body))
    story.append(bullet_list(STATISTICS))

    # -- 7. screenshots -----------------------------------------------------
    story.append(P("7. Screenshots", h1))
    for name, caption in SCREENSHOTS:
        img_path = os.path.join(SHOTS, name)
        if not os.path.exists(img_path):
            continue
        width = 12.5 * cm if name != "theme-selector.png" else 4.2 * cm
        story.append(KeepTogether([Image(img_path, width=width,
                                         height=width * 740.0 / 1180.0 if name != "theme-selector.png"
                                         else width * 185.0 / 194.0),
                                   Spacer(1, 0.15 * cm),
                                   P(caption, small),
                                   Spacer(1, 0.4 * cm)]))

    story.append(Spacer(1, 0.5 * cm))
    story.append(P("Everything printed in dark red between square brackets is a placeholder and "
                   "has to be replaced with the real name, number or date before submission.", small))

    doc = SimpleDocTemplate(path, pagesize=A4, title=TITLE, author="CanvasKanban team",
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm)

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(2 * cm, 1.1 * cm, "CanvasKanban - Software Project Report")
        canvas.drawRightString(A4[0] - 2 * cm, 1.1 * cm, "Page %d" % canvas.getPageNumber())
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


# ---------------------------------------------------------------------------
# 3. The Word document
# ---------------------------------------------------------------------------

def build_docx(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Cm, Pt

    doc = Document()

    for style_name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 13)):
        try:
            doc.styles[style_name].font.size = Pt(size)
        except KeyError:
            pass

    def runs(paragraph, text):
        """Adds text, marking the [to be filled in] parts in yellow."""
        for part in PLACEHOLDER.split(text):
            if not part:
                continue
            run = paragraph.add_run(re.sub(r"</?[bi]>", "", part))
            if PLACEHOLDER.fullmatch(part):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                run.bold = True

    def heading(text, level=1):
        h = doc.add_heading("", level=level)
        runs(h, text)
        return h

    def para(text="", style=None, bold=False):
        p = doc.add_paragraph(style=style)
        runs(p, text)
        if bold:
            for r in p.runs:
                r.bold = True
        return p

    def bullets(items):
        for item in items:
            p = doc.add_paragraph(style="List Bullet")
            runs(p, item)

    def table(rows, widths=None, header=True):
        t = doc.add_table(rows=0, cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            for c, value in enumerate(row):
                p = cells[c].paragraphs[0]
                runs(p, str(value))
                if i == 0 and header:
                    for r in p.runs:
                        r.bold = True
        if widths:
            for row in t.rows:
                for c, w in enumerate(widths):
                    row.cells[c].width = Cm(w)
        return t

    # title page
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(t, TITLE)
    for r in t.runs:
        r.bold = True
        r.font.size = Pt(24)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(s, SUBTITLE)
    doc.add_paragraph()
    table([(k, v) for k, v in SUBMITTED], [5, 11])
    doc.add_paragraph()
    doc.add_heading("Submitted by", level=2)
    table([("S. No.", "Full name", "Enrollment number", "Role")] + TEAM, [1.5, 6, 5, 4])
    doc.add_page_break()

    doc.add_heading("1. Project information", level=1)
    doc.add_heading("1.1 Project title", level=2)
    para("CanvasKanban - A Desktop Kanban Board Application")
    doc.add_heading("1.2 Project description", level=2)
    para(DESCRIPTION)
    doc.add_heading("1.3 Project objectives", level=2)
    bullets(OBJECTIVES)
    doc.add_heading("1.4 Technologies used", level=2)
    table(TECHNOLOGIES, [3.5, 5, 7.5])
    doc.add_heading("1.5 Key features and functionalities", level=2)
    bullets(FEATURES)
    doc.add_heading("1.6 GitHub repository", level=2)
    para(REPOSITORY)
    para(REPOSITORY_NOTE)

    doc.add_heading("2. Team members and responsibilities", level=1)
    para("The team has five members. One member acts as the team leader and is responsible for "
         "planning, architecture, integration and delivery; the other four members own one module "
         "each. Every member wrote code, tests and documentation, and everybody reviewed the work "
         "of at least one other member.")
    table([("S. No.", "Full name", "Enrollment number", "Role")] + TEAM, [1.5, 6, 5, 4])

    doc.add_heading("3. Individual contributions", level=1)
    for i, m in enumerate(CONTRIBUTIONS, start=1):
        heading("3.%d %s" % (i, m["name"]), level=2)
        para("Enrollment number: %s" % m["enroll"])
        para("Role in the project: %s" % m["role"])
        para("Specific tasks completed by the member:", bold=True)
        bullets(m["tasks"])
        para("Features / modules developed by the member:", bold=True)
        bullets(m["modules"])
        para("Technologies and tools used by the member: %s" % m["tech"])

    doc.add_page_break()
    doc.add_heading("4. Contributions shared by two or more members", level=1)
    bullets(SHARED)

    doc.add_heading("5. Contribution matrix", level=1)
    para("'Owner' means the member designed and implemented that part; 'Support' means the "
         "member contributed to it, for example through pair work, testing or review.")
    table([MATRIX_HEADER] + [list(r) for r in MATRIX], [5.2, 2.3, 2.3, 2.3, 2.3, 2.3])

    doc.add_heading("6. Project statistics and evidence", level=1)
    para("The numbers below were taken from the repository and from the continuous integration "
         "runs, so they can be checked by anyone with access to it.")
    bullets(STATISTICS)

    doc.add_heading("7. Screenshots", level=1)
    for name, caption in SCREENSHOTS:
        img = os.path.join(SHOTS, name)
        if not os.path.exists(img):
            continue
        doc.add_picture(img, width=Cm(15) if name != "theme-selector.png" else Cm(5))
        cap = para(caption)
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)

    para("Everything highlighted in yellow is a placeholder and has to be replaced with the real "
         "name, number or date before submission.")

    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 4. Run
# ---------------------------------------------------------------------------

def main():
    pdf = build_pdf(os.path.join(HERE, "CanvasKanban-Project-Report.pdf"))
    docx = build_docx(os.path.join(HERE, "CanvasKanban-Project-Report.docx"))

    placeholders = set()
    for value in (TITLE, DESCRIPTION, REPOSITORY) + tuple(x for pair in SUBMITTED for x in pair):
        placeholders.update(PLACEHOLDER.findall(value))
    for group in (OBJECTIVES, FEATURES, STATISTICS, SHARED):
        for value in group:
            placeholders.update(PLACEHOLDER.findall(value))
    for member in CONTRIBUTIONS:
        for value in [member["name"], member["enroll"]] + member["tasks"] + member["modules"]:
            placeholders.update(PLACEHOLDER.findall(value))
    for row in TEAM:
        for value in row:
            placeholders.update(PLACEHOLDER.findall(value))

    print("written:")
    print("  " + pdf)
    print("  " + docx)
    print("\nplaceholders still in the document (%d):" % len(placeholders))
    for value in sorted(placeholders):
        print("  " + value)


if __name__ == "__main__":
    main()
