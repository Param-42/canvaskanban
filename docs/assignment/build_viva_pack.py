#!/usr/bin/env python3
"""
Builds the two hand-out documents for the CanvasKanban assignment.

    1. "CanvasKanban - Simplified Project Report"      (short version of the report)
    2. "CanvasKanban - Viva Speaking Notes"            (what each member says)

Both come out as a PDF and as an editable Word file. Everything that still has
to be filled in (course, college, faculty, the five names and the five
enrollment numbers) is written as a [bracketed placeholder]. In the PDF the
placeholders are printed in dark red, in the Word file they are also
highlighted, so they are impossible to miss. They are all defined once, in the
DETAILS and TEAM tables right below.

Run it with the same interpreter that has reportlab and python-docx installed:

    /tmp/opencode/pdfenv/bin/python docs/assignment/build_viva_pack.py

The generated files are listed in .gitignore, because the filled-in versions
contain personal details. This script and the numbers in it are the source.
"""

from __future__ import annotations

import os
import re
import shutil
import struct
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
SCREENSHOTS = REPO / "docs" / "screenshots"

REPO_URL = "https://github.com/Param-42/canvaskanban"

APP_NAME = "CanvasKanban"
REPORT_TITLE = "CanvasKanban - A Desktop Kanban Board Application"

# ---------------------------------------------------------------------------
# 1. The details that have to be filled in by hand
# ---------------------------------------------------------------------------

DETAILS = {
    "course": "[Course name]",
    "college": "[College or university name]",
    "department": "[Department]",
    "faculty": "[Faculty name]",
    "year": "2026 - 2027",
    "date": "[Date of submission]",
}

# name, enrollment number, role
TEAM = [
    ("[Member 1 - Full Name]", "[Enrollment Number]",
     "Team leader - build, packaging and integration"),
    ("[Member 2 - Full Name]", "[Enrollment Number]",
     "Database and local storage"),
    ("[Member 3 - Full Name]", "[Enrollment Number]",
     "Data access objects and business rules"),
    ("[Member 4 - Full Name]", "[Enrollment Number]",
     "JavaFX user interface"),
    ("[Member 5 - Full Name]", "[Enrollment Number]",
     "Drag and drop, themes and documentation"),
]

# ---------------------------------------------------------------------------
# 2. Small helpers that describe content once and render it twice
#
# A document is a list of blocks. Every block is a tuple:
#   ("h1"|"h2"|"h3", text)
#   ("p", text)                      a normal paragraph
#   ("small", text)                  smaller grey note
#   ("bullets", [text, ...])
#   ("steps", [text, ...])           numbered list
#   ("table", rows, weights)         rows[0] is the header row
#   ("tableb", rows, weights)        same, but the first column is bold
#   ("code", text)                   monospace box, shorthand kept as it is
#   ("say", title, text)             a highlighted "say this out loud" box
#   ("diagram",)                     the layer picture
#   ("images", [(file, caption), ...])
#   ("pagebreak",)
# ---------------------------------------------------------------------------

PLACEHOLDER = re.compile(r"\[([^\[\]]+)\]")


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def image_size(path):
    """The pixel size of a PNG, read from its header. None for other formats."""
    with open(path, "rb") as handle:
        header = handle.read(24)
    if header[:8] != b"\x89PNG\r\n\x1a\n":
        return None
    return struct.unpack(">II", header[16:24])


def display_width(path, max_width_pt):
    """Shrinks small pictures: a screenshot is never shown larger than 150 dpi."""
    size = image_size(path)
    if not size:
        return max_width_pt
    return min(max_width_pt, size[0] * 72.0 / 150.0)


def _pdf_text(text: str) -> str:
    """Escapes the text and paints [placeholders] in dark red."""
    return PLACEHOLDER.sub(r'<font color="#B00020"><b>[\1]</b></font>', _escape(text))


# ---------------------------------------------------------------------------
# 3. The content of the simplified report
# ---------------------------------------------------------------------------

REPORT = [
    ("h1", "How to use this short report"),
    ("p", "This is the short version of the full project report. It is written so that it "
          "can be read out loud: every section is one idea, and the tables can be pointed at "
          "while speaking. The companion document, \"CanvasKanban - Viva Speaking Notes\", "
          "has the per-member speaking notes and the longer demo script."),
    ("bullets", [
        "Start with section 1 (what the project is) and section 2 (how the four layers fit together).",
        "Show the live application while reading section 7, the demo order.",
        "Use section 8 to show all six themes without clicking anything.",
        "Everything printed in dark red is a detail that still has to be filled in.",
    ]),

    ("h1", "1. Project information"),
    ("h2", "1.1 Project title"),
    ("p", "CanvasKanban - A Desktop Kanban Board Application"),
    ("h2", "1.2 Project description"),
    ("p", "CanvasKanban is a desktop application for one person to organise tasks. It is "
          "written in Java 21 with a JavaFX interface. The board has three columns - To Do, "
          "In Progress and Done - and a task can be moved between them by dragging the card."),
    ("p", "A task has a title, an optional description, a priority (Low, Medium or High) and "
          "an optional due date. Every board can use one of six colour themes, and the theme "
          "and the last opened board are remembered the next time the program starts."),
    ("p", "All boards, tasks and settings are stored in a single SQLite file in the user's own "
          "data folder. The application does not need an account, a server or an internet "
          "connection, and it can be installed like any other desktop program on Windows or "
          "Linux without installing Java first."),
    ("h2", "1.3 Project objectives"),
    ("bullets", [
        "Build a complete desktop application with the code separated into clear layers.",
        "Learn how to store and update information in a relational database using SQL.",
        "Create an interactive JavaFX interface with task dialogs, drag and drop and themes.",
        "Use Maven, automated tests and GitHub Actions to build and check the application.",
        "Provide installers so that users can install the application without setting up Java.",
    ]),

    ("h2", "1.4 Technologies used"),
    ("table", [
        ("Area", "Technology", "Simple purpose"),
        ("Programming", "Java 21", "The language the application is written in."),
        ("User interface", "JavaFX 21", "Creates the windows, the board columns, the task cards and the dialogs."),
        ("Appearance", "JavaFX CSS", "Controls colours, borders and spacing; one file per theme."),
        ("Database", "SQLite with JDBC", "Stores boards, tasks and settings in one local file."),
        ("Application logic", "Java services and DAOs", "Applies the rules and reads or writes the database."),
        ("Build", "Maven 3.9", "Compiles the code and manages the libraries and the tests."),
        ("Testing", "JUnit 5", "Runs the automated checks on the application."),
        ("Packaging", "jpackage and jlink", "Creates the installers with their own Java runtime."),
        ("Automation", "GitHub Actions", "Builds and tests the project after every push."),
        ("Install scripts", "Bash and PowerShell", "Install, update or remove the application."),
        ("Version control", "Git and GitHub", "Stores the code and publishes the release files."),
    ], [2.1, 2.0, 5.5]),

    ("h1", "2. How the application works"),
    ("p", "The program is split into four layers. Each layer has one job and only talks to the "
          "layer below it, which keeps a change in the database from forcing a change in the "
          "screen, and makes the rules testable without opening a window."),
    ("table", [
        ("Layer", "What it does"),
        ("1. User interface (UI)", "Shows the board and the cards and collects the actions of the user."),
        ("2. Service layer", "Checks the input and applies the rules, for example where a new task is placed."),
        ("3. Data access objects (DAOs)", "Run the SQL commands that create, read, update and delete records."),
        ("4. SQLite database", "Stores boards, tasks and settings in one file on the user's computer."),
    ], [2.6, 7.0]),
    ("h2", "2.2 The layers in one picture"),
    ("diagram",),
    ("h2", "2.3 What happens when a task is created"),
    ("steps", [
        "The user clicks \"New task\" in the toolbar and fills in the form.",
        "MainView asks BoardService to create the task - it never writes SQL itself.",
        "BoardService builds a Task object, sets the creation date and calls TaskDao.",
        "TaskDao writes the task with a prepared statement and reads back the generated id.",
        "MainView reloads the three columns, so the new card appears immediately.",
    ]),
    ("small", "Editing, deleting, moving and changing the theme all follow this same path: "
              "the interface calls the service, and only the DAO layer talks to the database."),

    ("h1", "3. Main features"),
    ("bullets", [
        "Boards: create, rename and delete boards; a fresh installation already contains one example board.",
        "Tasks: title, description, priority (Low, Medium, High), creation date and an optional due date.",
        "Three fixed columns: To Do, In Progress and Done.",
        "Drag and drop: drag a card to another column to change its status, or up and down to change its order; a blue line shows where it will land.",
        "Task menu: right-click a card to read the full description, edit it, or delete it after a confirmation.",
        "Six themes: Minimal Light, Dark Mode, Midnight Blue, Forest Green, Lavender and Warm Sand.",
        "Saved preferences: the last opened board and each board's theme are remembered after a restart.",
        "Overdue check: a task that is past its due date and not finished yet is marked as overdue.",
        "Status bar: shows the last action, the database location and the number of tasks on the board.",
        "Local storage: no account, no server and no internet connection are needed.",
        "Installers: Windows MSI and ZIP, Linux DEB, RPM and tar.gz, plus a plain JAR for anyone who already has Java.",
        "Installer scripts: install.sh for Linux and install.ps1 for Windows, which install, update and remove the application.",
    ]),
    ("h2", "3.1 What to point at while the examiner looks at the screen"),
    ("table", [
        ("Feature", "Where to point", "What to say"),
        ("Boards", "The board list and the board buttons in the toolbar",
         "Boards are kept in the database, so a second board is another set of tasks."),
        ("Tasks", "A card, and the \"New task\" button",
         "A task is one row in the tasks table; the card is only the way it is shown."),
        ("Priority and due date", "The coloured chip and the date on a card",
         "Both are columns of the same row, and an overdue task is marked when it is loaded."),
        ("Drag and drop", "Drag a card to another column",
         "The drop position becomes the new status and the new position, saved inside one transaction."),
        ("Right-click menu", "Right-click a card",
         "The full description, editing and deleting are all handled by the same service."),
        ("Themes", "The theme selector in the toolbar",
         "The theme is stored with the board, so opening that board again restores it."),
        ("Status bar", "The bottom of the window",
         "It shows the last action and the exact location of the database file."),
        ("Restart", "Close and start the application again",
         "The board, the tasks and the theme come back, because they were never kept only in memory."),
    ], [1.7, 2.9, 5.0]),

    ("h2", "3.2 Useful terms"),
    ("table", [
        ("Term", "Meaning in this project"),
        ("DAO", "Data Access Object: a class that contains the database operations for one table."),
        ("JDBC", "The Java API used to connect Java code to a database."),
        ("Prepared statement", "An SQL command with placeholders for the values, so the SQL text and the data stay separate."),
        ("Transaction", "A group of database changes that are saved together or not at all."),
        ("SQLite", "A database that lives in a single file instead of a server."),
        ("Maven", "The build tool that compiles the project and manages the libraries."),
        ("CI (continuous integration)", "Building and testing automatically after every push to GitHub."),
        ("jpackage / jlink", "JDK tools that build an installer containing its own Java runtime."),
    ], [3.0, 6.6]),

    ("h1", "4. Team members and responsibilities"),
    ("p", "The project is a five-member effort. The team leader managed the planning, the "
          "architecture and the delivery. Each of the other four members led one area. "
          "Everyone also helped with testing, reviews and integration. The names and "
          "enrollment numbers are on the cover page; the member numbers below are the same "
          "role labels that are used throughout the report."),
    ("tableb", [
        ("Member", "Main role", "What the member did"),
        ("Member 1", "Team leader - build and delivery",
         "Created the repository and the Maven setup, planned the code structure, combined the "
         "modules, prepared the Windows and Linux packages and the installer scripts, and "
         "configured GitHub Actions and the release files."),
        ("Member 2", "Database and storage",
         "Designed the SQLite tables, the index and the relationships, wrote the connection code, "
         "used prepared statements, handled the per-system data folder and the date conversion, "
         "and wrote the database tests."),
        ("Member 3", "Data access and business rules",
         "Wrote the DAO classes and BoardService, added the create, read, update and delete "
         "operations, handled moving and ordering a task inside a transaction, and created the "
         "model classes and the service tests."),
        ("Member 4", "JavaFX user interface",
         "Built the main screen, the toolbar, the three columns, the status bar, the task cards "
         "and the create, edit and delete dialogs, connected the interface to the service layer "
         "and helped test the interface."),
        ("Member 5", "Drag and drop, themes and documentation",
         "Built the drag-and-drop handling and the drop indicator, created the six CSS themes and "
         "the theme selector, saved the theme per board, wrote the user and developer guides and "
         "took the screenshots and the theme tests."),
    ], [1.3, 2.6, 6.0]),
    ("h2", "4.1 The main files of each member"),
    ("tableb", [
        ("Member", "Examples of files and modules"),
        ("Member 1", "pom.xml, package.sh, package.bat, .github/workflows/build.yml, install.sh, install.ps1"),
        ("Member 2", "AppPaths.java, Database.java, Schema.java, SqlDates.java, DataAccessException.java"),
        ("Member 3", "BoardDao.java, TaskDao.java, SettingDao.java, BoardService.java, the Board, Task, Priority, TaskStatus and Theme classes"),
        ("Member 4", "Launcher.java, MainApp.java, MainView.java, Dialogs.java, TaskCardView.java"),
        ("Member 5", "BoardColumnView.java, ThemeManager.java, ThemeListCell.java, base.css and the six theme files, README and docs/"),
    ], [1.3, 8.3]),

    ("h1", "5. Work shared by the team"),
    ("p", "Some work cannot be divided, so it was done in pairs and then reviewed by everybody:"),
    ("bullets", [
        "Members 3 and 4 tested complete board workflows together and found the task-ordering problems that appeared while dragging.",
        "Members 4 and 5 checked that the drop highlight and the blue position line stayed visible in all six themes.",
        "Members 2 and 3 agreed on the table names, the column types and the ordering rules, and then reviewed each other's SQL.",
        "Members 1 and 3 joined the service layer to the build and the tests.",
        "All five members installed a packaged version, tried it and reviewed the documentation before submission.",
    ]),
    ("h2", "5.1 Contribution summary"),
    ("p", "\"Owner\" is the member who mainly implemented an area; \"support\" is the member "
          "who helped with coding, testing or review."),
    ("table", [
        ("Area", "Owner", "Support"),
        ("Planning, architecture and integration", "Member 1", "Members 2, 3, 4 and 5"),
        ("Maven build configuration", "Member 1", "-"),
        ("SQLite schema, connections and data paths", "Member 2", "Members 1 and 3"),
        ("DAO layer and business rules", "Member 3", "Member 2"),
        ("JavaFX interface, dialogs and task cards", "Member 4", "Members 1, 3 and 5"),
        ("Drag and drop and task ordering", "Member 5", "Members 3 and 4"),
        ("Themes and CSS styling", "Member 5", "Member 4"),
        ("Automated testing", "Members 2, 3, 4 and 5", "Member 1"),
        ("Packaging and release files", "Member 1", "Member 5"),
        ("Continuous integration", "Member 1", "Members 3 and 5"),
        ("Documentation and screenshots", "Member 5", "Members 1, 2, 3 and 4"),
    ], [4.2, 2.4, 2.6]),

    ("h1", "6. Project statistics and testing"),
    ("table", [
        ("Item", "Details taken from the project"),
        ("Application code", "22 Java files and 2,627 lines of code."),
        ("Automated tests", "7 test classes, 1,255 lines of test code and 67 test methods."),
        ("What is tested", "The database layer, the DAO classes, the service rules, the theme files and the real JavaFX interface."),
        ("Styling", "7 CSS files and 664 lines: one base layout file and six theme files."),
        ("Database", "3 tables (boards, tasks, settings) plus one index for the column lookups, in a single SQLite file."),
        ("Automatic builds", "Every push runs \"mvn clean package\" and the tests on Ubuntu and Windows; the interface tests run on a virtual screen (Xvfb)."),
        ("Release", "Version 1.0.0 publishes six files: Windows MSI and ZIP, Linux DEB, RPM and tar.gz, and a plain JAR."),
        ("Package checks", "The DEB, RPM, tar.gz and MSI are built automatically on matching systems: the DEB and the tar.gz on Ubuntu, the RPM on Fedora and the MSI with the WiX toolset on Windows. The Linux tar.gz package was installed and started on a Fedora machine while the project was being developed."),
    ], [2.6, 7.0]),
    ("h2", "6.1 GitHub repository"),
    ("p", REPO_URL),
    ("p", "The repository holds the source code, the build and packaging scripts, the GitHub "
          "Actions workflow, the screenshots and the user and developer documentation."),

    ("h1", "7. Demo order (what to click, in order)"),
    ("p", "A demonstration of about five minutes. Each step says what to do and what to say while doing it."),
    ("steps", [
        "Start the application - \"It is a normal desktop program. There is no server and the data stays on this computer.\"",
        "Point at the board and the three columns - \"A board always has these three columns; the cards are the tasks.\"",
        "Click \"New task\", fill it in and save - \"This goes through the service layer and into the SQLite file.\"",
        "Drag that card to Done - \"The drag changes the status and the position, and a blue line shows where it will land.\"",
        "Right-click a card - \"The menu shows the full description, editing and deleting with a confirmation.\"",
        "Change the theme in the toolbar - \"The theme belongs to the board, and base.css only holds the layout.\"",
        "Point at the status bar - \"This shows the last action and the exact database file on this machine.\"",
        "Close the program, start it again - \"The board, the tasks and the theme are all back, because they were saved, not only shown.\"",
    ]),

    ("h1", "8. Screenshots"),
    ("p", "The same board in the six themes, the theme selector, and a card after it was moved."),
    ("images", [
        ("minimal-light.png", "Minimal Light, the default theme: three columns and cards with a priority colour and a due date."),
        ("dark-mode.png", "Dark Mode: the theme is stored with the board and restored when it is opened again."),
        ("midnight-blue.png", "Midnight Blue."),
        ("forest-green.png", "Forest Green, after a task was added, edited, moved and deleted."),
        ("lavender.png", "Lavender."),
        ("warm-sand.png", "Warm Sand."),
        ("theme-selector.png", "The theme selector in the toolbar."),
        ("drag-and-drop-result.png", "After a drag: the card is in its new column and the new order is saved."),
    ]),

    ("h1", "9. Conclusion"),
    ("p", "CanvasKanban is a local desktop task-management application. It combines a JavaFX "
          "interface, a SQLite database, drag-and-drop task management, six themes, automated "
          "tests and installers for Windows and Linux. Its layered structure keeps the "
          "interface, the application rules and the database work apart, which is what makes "
          "the project easy to extend, to test and to explain."),
]

REPORT_SUBTITLE = "Simplified Software Project Report"


# ---------------------------------------------------------------------------
# 4. The content of the viva speaking notes
# ---------------------------------------------------------------------------

NOTES = [
    ("h1", "How to use these notes"),
    ("p", "These notes are written to be read out loud. Part 1 is for the whole team - the "
          "short introduction, the picture of the layers, the demo and the list of who answers "
          "what. Part 2 is one section per member: the files to open, the sentences to say, "
          "the real code to point at, and the questions that are most likely to come."),
    ("bullets", [
        "Each member reads their own part of Part 2 twice, and reads Part 1 once.",
        "Every code box is real code from the repository; lines left out are marked with // ... or # ... Open the file named above the box and point at it on the screen.",
        "The coloured boxes are the sentences to say. Everything else is background knowledge.",
        "Member 1 to Member 5 are the role labels used in the report; write your own name in the box on the cover.",
    ]),

    ("h1", "Part 1 - What everybody should know"),
    ("say", "If you have only ten seconds",
     "CanvasKanban is a desktop Kanban board written in Java 21 with a JavaFX interface. "
     "Boards, tasks and settings are kept in one local SQLite file, so the program needs no "
     "server and no internet connection. The interface, the rules and the database work are "
     "kept in separate layers, and ready-made installers are published for Windows and Linux."),
    ("h2", "1.2 How the code fits together"),
    ("p", "Say the classes in this order and the whole program makes sense:"),
    ("table", [
        ("Class or layer", "Its one job"),
        ("MainApp", "Starts JavaFX, prepares the database and the service, then builds the window."),
        ("MainView", "The screen: toolbar, three columns, task cards, dialogs and the status bar."),
        ("BoardService", "The rules of the application, and the only thing the screen is allowed to call."),
        ("BoardDao, TaskDao, SettingDao", "The SQL: one class per table, using prepared statements."),
        ("Database", "Opens the SQLite connections and switches foreign keys on."),
        ("canvaskanban.db", "The SQLite file in the user's data folder where everything is stored."),
    ], [3.0, 6.6]),
    ("diagram",),
    ("h2", "1.3 The five members at a glance"),
    ("tableb", [
        ("Member", "Area", "Files you should be able to open"),
        ("Member 1", "Build, packaging and delivery", "pom.xml, package.sh, package.bat, .github/workflows/build.yml, install.sh, install.ps1"),
        ("Member 2", "Database and storage", "Schema.java, Database.java, AppPaths.java, SqlDates.java"),
        ("Member 3", "DAOs and business rules", "BoardService.java, BoardDao.java, TaskDao.java, SettingDao.java, the model classes"),
        ("Member 4", "JavaFX interface", "MainApp.java, MainView.java, Dialogs.java, TaskCardView.java, Launcher.java"),
        ("Member 5", "Drag and drop, themes, docs", "BoardColumnView.java, ThemeManager.java, base.css and the six theme files, README and docs/"),
    ], [1.3, 2.6, 6.0]),
    ("h2", "1.4 The demo script (about five minutes)"),
    ("p", "Before the demonstration, open these three things: the repository in a browser "
          "(the Releases page and the Actions tab), the project in the IDE, and a terminal."),
    ("steps", [
        "Run the application. Say: \"It is a single desktop program; the data is a file on this computer.\"",
        "Show the three columns and the existing example board. Say: \"A new installation creates one example board so nothing looks empty.\"",
        "Create a task in In Progress. Say the flow out loud: \"The view calls the service, the service builds the object, and the task DAO inserts the row.\"",
        "Drag the task to Done. Say: \"The new status and the new position are written together in one transaction, so a move cannot end up half saved.\"",
        "Right-click a card and show the description, then cancel the delete dialog. Say: \"Deleting asks for confirmation first.\"",
        "Point at the status bar. Say: \"This is the exact SQLite file. Boards and tasks are rows in that file.\"",
        "Switch the theme to Dark Mode, then back. Say: \"The layout is in one CSS file and each theme only changes the colours.\"",
        "Close and restart the program. Say: \"The board, the tasks and the theme all come back, and the last opened board is remembered.\"",
    ]),
    ("h2", "1.5 Who answers what"),
    ("table", [
        ("If the examiner asks about ...", "The member who answers"),
        ("Maven, libraries, packaging, installers, GitHub Actions, releases", "Member 1"),
        ("Tables, columns, keys, connections, where the database file is kept", "Member 2"),
        ("Service rules, DAOs, SQL, ordering, transactions, model classes", "Member 3"),
        ("Startup, the screen, dialogs, cards, how the view talks to the service", "Member 4"),
        ("Dragging a card, the drop position, themes, CSS, screenshots, documentation", "Member 5"),
        ("\"Why JavaFX?\", \"Why SQLite?\", \"What is Maven?\"", "Any member - the answers are in Part 3"),
        ("\"Which part did you do?\"", "Each member answers for their own area"),
        ("\"How did you work as a team?\"", "Member 1 starts, then the others add what they worked on together"),
    ], [5.6, 4.0]),
    ("h2", "1.6 Prove it, don't only say it"),
    ("p", "These four things turn a claim into evidence. Have them ready:"),
    ("bullets", [
        "The test run: \"mvn test\" finishes with 67 tests and a green result.",
        "The database: show the row behind a task you just created, either with the sqlite3 command below or in a SQL client - DBeaver, SQL Workbench/J or any JDBC client, the settings are in docs/DATABASE.md.",
        "The release page: six downloadable files for Windows and Linux.",
        "The runs page: the same build and the same tests run automatically on Ubuntu and Windows.",
    ]),
    ("code", "mvn test                                   # 67 tests\n"
             "mvn javafx:run                             # start the application\n"
             "sqlite3 ~/.local/share/canvaskanban/canvaskanban.db \\\n"
             "  \"SELECT b.name, t.status, t.title FROM boards b \" \\\n"
             "  \"JOIN tasks t ON t.board_id = b.id ORDER BY t.id;\"\n"
             "\n"
             "# the same database in any SQL client, over JDBC:\n"
             "#   driver  org.sqlite.JDBC\n"
             "#   url     jdbc:sqlite:/home/you/.local/share/canvaskanban/canvaskanban.db\n"
             "#   read only: add ?mode=ro, or close the application before writing"),

    ("h1", "Part 2 - Member by member"),

    # ------------------------------------------------------------------ M1
    ("h2", "Member 1 - Build, packaging and project integration"),
    ("p", "Open these files: pom.xml, package.sh, package.bat, .github/workflows/build.yml, install.sh, install.ps1."),
    ("say", "Say this first",
     "My part is the project setup and the delivery. I configured Maven for Java 21 and the "
     "libraries we use, wrote the scripts that build an installer for Windows and for Linux "
     "with the Java runtime included, and set up GitHub Actions so that every push is compiled "
     "and tested automatically and a tagged release publishes the six downloadable files."),
    ("h3", "1. The versions the whole project uses (pom.xml)"),
    ("code", "<properties>\n"
             "    <project.build.sourceEncoding>UTF-8</project.build.sourceEncoding>\n"
             "    <maven.compiler.release>21</maven.compiler.release>\n"
             "\n"
             "    <javafx.version>21.0.12</javafx.version>\n"
             "    <sqlite.version>3.53.4.0</sqlite.version>\n"
             "    <slf4j.version>2.0.20</slf4j.version>\n"
             "    <junit.version>5.14.4</junit.version>\n"
             "\n"
             "    <main.class>com.canvaskanban.MainApp</main.class>\n"
             "    <launcher.class>com.canvaskanban.Launcher</launcher.class>\n"
             "</properties>"),
    ("p", "What it means: pom.xml is Maven's description of the project. It fixes the Java "
          "version and every library version here, so the same versions are used on every "
          "machine and in the build server."),
    ("h3", "2. Building an installer (package.sh)"),
    ("code", "APP_NAME=\"CanvasKanban\"\n"
             "MAIN_CLASS=\"com.canvaskanban.Launcher\"     # a plain class, never MainApp\n"
             "MAIN_JAR=\"${ARTIFACT}-${VERSION}.jar\"\n"
             "\n"
             "jpackage \\\n"
             "  --type \"$type\" \\\n"
             "  --name \"$APP_NAME\" \\\n"
             "  --app-version \"$VERSION\" \\\n"
             "  --input \"$LIBS_DIR\" \\\n"
             "  --main-jar \"$MAIN_JAR\" \\\n"
             "  --main-class \"$MAIN_CLASS\" \\\n"
             "  --module-path \"$JAVAFX_DIR\" \\\n"
             "  --add-modules javafx.controls,java.sql,java.logging,java.desktop \\\n"
             "  --java-options \"-Dfile.encoding=UTF-8\" \\\n"
             "  --dest \"$DEST_DIR\""),
    ("p", "What it means: \"mvn clean package\" compiles the code, runs the tests and creates "
          "the JAR. jpackage then wraps the JAR and its libraries into a real application "
          "package together with a small Java runtime, which is why the user does not have to "
          "install Java. The same script produces a .deb and a .rpm on Linux and, with "
          "package.bat, an MSI and a ZIP on Windows."),
    ("h3", "3. Building and testing automatically (build.yml)"),
    ("code", "on:\n"
             "  push:\n"
             "    branches: [ main, master ]\n"
             "    tags: [ 'v*' ]\n"
             "  pull_request:\n"
             "\n"
             "jobs:\n"
             "  test-and-package:\n"
             "    runs-on: ${{ matrix.os }}\n"
             "    strategy:\n"
             "      matrix:\n"
             "        os: [ ubuntu-latest, windows-latest ]"),
    ("p", "What it means: the workflow runs on Linux and on Windows, because an installer can "
          "only be built on the system it is meant for. Each run compiles the project and runs "
          "all the tests; the interface tests run on a virtual screen so they work on a server "
          "without a monitor."),
    ("h3", "4. Release files with names that never change"),
    ("code", "for f in *; do\n"
             "  case \"$f\" in\n"
             "    *-linux-*.tar.gz) mv \"$f\" canvaskanban-linux-x86_64.tar.gz ;;\n"
             "    *.deb)            mv \"$f\" canvaskanban-linux-x86_64.deb ;;\n"
             "    *.rpm)            mv \"$f\" canvaskanban-linux-x86_64.rpm ;;\n"
             "    *.msi)            mv \"$f\" CanvasKanban-windows-x64.msi ;;\n"
             "    *.jar)            mv \"$f\" canvaskanban.jar ;;\n"
             "  esac\n"
             "done"),
    ("p", "What it means: the packaging tools put the version number into the file names, "
          "which changes with every release. Renaming them to fixed names means the download "
          "links in the README keep working for the next version as well. A second step "
          "removes any file on the release that is not one of the six expected names."),
    ("h3", "If the examiner asks ..."),
    ("table", [
        ("Question", "Your answer"),
        ("Why Maven and not a simple javac script?",
         "Maven records the library versions and gives one repeatable command that compiles, tests and packages - the same command on my laptop and on the build server."),
        ("Why are some packages built in separate jobs?",
         "A native installer can only be built on its own system: the .deb and the .rpm on Linux and the .msi on Windows, so each job runs on the matching machine."),
        ("What exactly does the user download?",
         "The DEB or RPM on Linux, the MSI or the ZIP on Windows, or the JAR if Java is already installed. The installers contain their own runtime, so Java is not needed."),
        ("What does install.sh do?",
         "It finds the right file for the machine, verifies it, installs it and can update or remove it again; install.ps1 does the same on Windows."),
        ("How do you know the project still builds?",
         "Every push starts the workflow and the result is visible in the Actions tab; the release is only published when those jobs succeed."),
    ], [3.4, 6.2]),
    ("h3", "Terms to know"),
    ("bullets", [
        "Maven - the build tool: it downloads libraries, compiles, tests and packages the project.",
        "Dependency - a library the project needs, such as JavaFX or the SQLite JDBC driver.",
        "jpackage - the JDK tool that turns a JAR into an installer with its own Java runtime.",
        "CI - continuous integration: building and testing automatically after every change.",
        "Artifact - a built file, for example the JAR or the .deb produced by a build.",
    ]),
    ("small", "Do not say that you wrote the whole project: your part is the build, the "
              "packaging and the integration of the other members' modules."),

    # ------------------------------------------------------------------ M2
    ("h2", "Member 2 - Database and local storage"),
    ("p", "Open these files: database/Schema.java, database/Database.java, database/AppPaths.java, database/SqlDates.java."),
    ("say", "Say this first",
     "My part is the storage. The application keeps everything in one SQLite file, and my code "
     "defines the three tables, opens the connections, switches on foreign keys and decides "
     "where the file is kept on Windows, Linux and macOS. Because the database is a local file, "
     "the program needs no server and works without an internet connection."),
    ("h3", "1. The tables (Schema.java)"),
    ("code", "CREATE TABLE IF NOT EXISTS boards (\n"
             "  id          INTEGER PRIMARY KEY AUTOINCREMENT,\n"
             "  name        TEXT    NOT NULL,\n"
             "  description TEXT    NOT NULL DEFAULT '',\n"
             "  theme       TEXT    NOT NULL DEFAULT 'minimal-light',\n"
             "  created_at  TEXT    NOT NULL\n"
             ")"),
    ("code", "CREATE TABLE IF NOT EXISTS tasks (\n"
             "  id          INTEGER PRIMARY KEY AUTOINCREMENT,\n"
             "  board_id    INTEGER NOT NULL,\n"
             "  title       TEXT    NOT NULL,\n"
             "  description TEXT    NOT NULL DEFAULT '',\n"
             "  priority    TEXT    NOT NULL DEFAULT 'MEDIUM',\n"
             "  status      TEXT    NOT NULL DEFAULT 'TODO',\n"
             "  position    INTEGER NOT NULL DEFAULT 0,\n"
             "  created_at  TEXT    NOT NULL,\n"
             "  due_date    TEXT,\n"
             "  FOREIGN KEY (board_id) REFERENCES boards (id) ON DELETE CASCADE\n"
             ")"),
    ("p", "What it means: three tables - boards, tasks and settings - are created when the "
          "program starts for the first time, if they are not there yet. The tasks table keeps "
          "board_id as a foreign key, and ON DELETE CASCADE means that deleting a board also "
          "deletes its tasks, so no task can be left behind without a board. An index on the "
          "board and status columns keeps loading a column fast."),
    ("h3", "2. Opening the database (Database.java)"),
    ("code", "public Connection getConnection() throws SQLException {\n"
             "    Connection connection = DriverManager.getConnection(jdbcUrl);\n"
             "    try (Statement statement = connection.createStatement()) {\n"
             "        // SQLite disables foreign keys by default, so switch them on per connection.\n"
             "        statement.execute(\"PRAGMA foreign_keys = ON\");\n"
             "    }\n"
             "    return connection;\n"
             "}"),
    ("p", "What it means: the JDBC URL \"jdbc:sqlite:<path>\" says which file to open. Every "
          "operation gets its own connection and closes it again with try-with-resources, so "
          "there is no long-lived open handle. Foreign keys are switched on for each connection, "
          "because SQLite has them off by default."),
    ("h3", "3. Where the file is kept (AppPaths.java)"),
    ("code", "if (osName.contains(\"win\")) {\n"
             "    String appData = System.getenv(\"APPDATA\");\n"
             "    // ... %APPDATA%\\CanvasKanban\n"
             "}\n"
             "if (osName.contains(\"mac\")) {\n"
             "    return Paths.get(userHome, \"Library\", \"Application Support\", APP_NAME)...\n"
             "}\n"
             "return Paths.get(userHome, \".local\", \"share\", APP_NAME_LOWER)...   // Linux"),
    ("p", "What it means: on Windows the file is in %APPDATA%\\CanvasKanban, on Linux in "
          "~/.local/share/canvaskanban and on macOS in ~/Library/Application Support/"
          "CanvasKanban. Keeping it in the user's data folder instead of beside the program "
          "means the tasks survive an update or a reinstall. A system property or environment "
          "variable can point the program somewhere else, which is used while testing."),
    ("h3", "If the examiner asks ..."),
    ("table", [
        ("Question", "Your answer"),
        ("Why SQLite and not MySQL?",
         "The application is for one person on one computer. SQLite is a single file, needs no server, works offline and is enough for thousands of tasks."),
        ("Which tables are there?",
         "boards, tasks and settings. A board has many tasks through board_id, and the settings table keeps things like the last opened board and the default theme."),
        ("Why a new connection for every operation?",
         "SQLite is used from the single JavaFX thread, so one connection per operation is simple and safe; try-with-resources closes it and nothing is left open."),
        ("What if the same fix is needed on another computer?",
         "The data folder is chosen by the operating system, so the program behaves the same on Windows, Linux and macOS without any configuration."),
        ("How is a date stored in a text column?",
         "The due date is stored as text in the ISO form yyyy-MM-dd, which sorts and compares correctly and is easy to read in a database browser."),
        ("Can I open the database with a normal SQL client?",
         "Yes. It is a plain SQLite file, so DBeaver, SQL Workbench/J or any JDBC client can open it with the driver org.sqlite.JDBC and the URL jdbc:sqlite:<path to the file> - the same driver the application uses."),
        ("What if the client and the application both write?",
         "Only one writer is allowed at a time, so the second one waits and then fails with \"database is locked\". Reading is always safe; for writing we close the application first, or work on a copy made with VACUUM INTO."),
    ], [3.4, 6.2]),
    ("h3", "Terms to know"),
    ("bullets", [
        "SQLite - a database that is one file, with no separate server process.",
        "JDBC - the Java API for talking to a database.",
        "Schema - the definition of the tables, columns, keys and relationships.",
        "Foreign key - a column that points at a row in another table; ON DELETE CASCADE keeps the two in step.",
        "PRAGMA - a command for SQLite settings, used here to switch foreign keys on.",
    ]),
    ("small", "Do not say the database is a server: it is a file, and that is the point."),

    # ------------------------------------------------------------------ M3
    ("h2", "Member 3 - Data access objects and business rules"),
    ("p", "Open these files: service/BoardService.java, dao/BoardDao.java, dao/TaskDao.java, "
          "dao/SettingDao.java and the model classes (Board, Task, Priority, TaskStatus, Theme)."),
    ("say", "Say this first",
     "My part is the application rules and the database operations. BoardService is the layer "
     "the interface calls, and it is where the rules live, for example where a new task is "
     "placed. The DAO classes contain the SQL, one class per table, and moving a card updates "
     "the order inside a transaction so a move can never be saved half-done. I also wrote the "
     "model classes that hold a board, a task and its priority, status and theme."),
    ("h3", "1. A rule with almost no code (BoardService.createTask)"),
    ("code", "public Task createTask(long boardId, String title, String description,\n"
             "                       Priority priority, TaskStatus status, LocalDate dueDate) {\n"
             "    Task task = new Task(0L, boardId, title, description, priority, status,\n"
             "            LocalDate.now(), dueDate, 0);\n"
             "    return taskDao.insert(task);\n"
             "}"),
    ("p", "What it means: the service receives plain values from the interface, builds a Task "
          "object, fills in the creation date itself so the screen cannot get it wrong, and "
          "asks the DAO to save it. The interface never writes SQL."),
    ("h3", "2. Saving with a prepared statement (TaskDao.insert)"),
    ("code", "String sql = \"INSERT INTO tasks (board_id, title, description, priority, status, position,\"\n"
             "        + \" created_at, due_date) VALUES (?, ?, ?, ?, ?, ?, ?, ?)\";\n"
             "try (Connection connection = database.getConnection()) {\n"
             "    int position = task.getPosition() > 0\n"
             "            ? task.getPosition()\n"
             "            : nextPosition(connection, task.getBoardId(), task.getStatus());\n"
             "\n"
             "    try (PreparedStatement statement =\n"
             "                 connection.prepareStatement(sql, Statement.RETURN_GENERATED_KEYS)) {\n"
             "        statement.setLong(1, task.getBoardId());\n"
             "        statement.setString(2, task.getTitle());\n"
             "        statement.setString(3, task.getDescription());\n"
             "        statement.setString(4, task.getPriority().name());\n"
             "        statement.setString(5, task.getStatus().name());\n"
             "        statement.setInt(6, position);\n"
             "        statement.setString(7, createdAt.toString());\n"
             "        statement.setString(8, SqlDates.toText(task.getDueDate()));\n"
             "        statement.executeUpdate();\n"
             "\n"
             "        try (ResultSet keys = statement.getGeneratedKeys()) {\n"
             "            if (keys.next()) {\n"
             "                task.setId(keys.getLong(1));\n"
             "            }\n"
             "        }\n"
             "    }\n"
             "    // ..."),
    ("p", "What it means: the question marks are filled in by setLong, setString and the other "
          "setters, so the SQL text and the data never mix. That is safer and clearer than "
          "building the command by joining text together, and it also avoids quoting mistakes "
          "when a title contains an apostrophe. RETURN_GENERATED_KEYS gives back the id that "
          "SQLite assigned, and the new task is placed at the end of its column."),
    ("h3", "3. Moving a card safely (TaskDao.moveTask)"),
    ("code", "connection = database.getConnection();\n"
             "connection.setAutoCommit(false);\n"
             "\n"
             "Task task = findById(connection, taskId).orElse(null);\n"
             "if (task == null) {\n"
             "    connection.rollback();\n"
             "    return false;\n"
             "}\n"
             "\n"
             "TaskStatus sourceStatus = task.getStatus();\n"
             "List<Long> targetIds = idsInColumn(connection, task.getBoardId(), targetStatus, taskId);\n"
             "int index = Math.max(0, Math.min(targetIndex, targetIds.size()));\n"
             "targetIds.add(index, taskId);\n"
             "\n"
             "if (sourceStatus != targetStatus) {\n"
             "    // Closing the gap in the old column keeps the order tidy.\n"
             "    List<Long> sourceIds = idsInColumn(connection, task.getBoardId(), sourceStatus, taskId);\n"
             "    writeOrder(connection, sourceStatus, sourceIds);\n"
             "}\n"
             "writeOrder(connection, targetStatus, targetIds);\n"
             "\n"
             "connection.commit();\n"
             "// ... rollbackQuietly(connection) in the catch block"),
    ("p", "What it means: the order of a column is not a hidden property, it is the position "
          "column of each row. A move rewrites the positions of the target column and, if the "
          "status changed, closes the gap in the old column as well. Both changes go into one "
          "transaction: either everything is saved, or - if something fails - everything is "
          "rolled back, so the stored order can never be left half updated."),
    ("h3", "4. A rule that lives in the model (Task.isOverdue)"),
    ("code", "public boolean isOverdue() {\n"
             "    return dueDate != null && status != TaskStatus.DONE && dueDate.isBefore(LocalDate.now());\n"
             "}"),
    ("p", "What it means: a task is overdue only if it has a due date, that date is before "
          "today, and the task is not finished. The rule is written once in the model, so the "
          "card, the test and any future screen all agree on what \"overdue\" means."),
    ("h3", "If the examiner asks ..."),
    ("table", [
        ("Question", "Your answer"),
        ("Why is there a service layer between the interface and the DAOs?",
         "So that the rules live in one place. The screen can be changed or tested without knowing SQL, and the rules can be tested without opening a window."),
        ("Why prepared statements?",
         "The SQL text and the values are kept separate, so quotes and apostrophes in a title cannot break the command, and the same command is reused."),
        ("Why is a transaction needed to move a card?",
         "The positions of two columns change in the same logical action; with a transaction the move either completes completely or not at all."),
        ("How is the order of the cards stored?",
         "As a position number on every task row. After a move the affected columns are rewritten in the new order, which also survives a restart."),
        ("What is a DAO?",
         "A Data Access Object: one class per table that contains the create, read, update and delete operations for it, and nothing else."),
    ], [3.4, 6.2]),
    ("h3", "Terms to know"),
    ("bullets", [
        "Service layer - the place for the application rules, between the screen and the SQL.",
        "DAO - Data Access Object, one class per table, holding the SQL for that table.",
        "CRUD - create, read, update, delete: the four operations every DAO offers.",
        "Transaction - a group of changes that are committed together or rolled back.",
        "Model class - a plain Java object that represents one row, for example Task or Board.",
    ]),
    ("small", "Say that the interface never runs SQL itself - that one sentence explains the whole design."),

    # ------------------------------------------------------------------ M4
    ("h2", "Member 4 - The JavaFX user interface"),
    ("p", "Open these files: Launcher.java, MainApp.java, ui/MainView.java, ui/Dialogs.java, ui/TaskCardView.java."),
    ("say", "Say this first",
     "My part is the screen. MainApp starts JavaFX, opens the database and creates the service "
     "before the window appears. MainView builds the toolbar, the three columns and the status "
     "bar, and TaskCardView with Dialogs handles clicking a card, opening the form and asking "
     "before deleting. The interface only calls BoardService - it never writes SQL."),
    ("h3", "1. Starting the application (MainApp.java)"),
    ("code", "@Override\n"
             "public void init() {\n"
             "    try {\n"
             "        database = new Database();\n"
             "        database.initialize();\n"
             "        service = new BoardService(database);\n"
             "        service.createSampleDataIfEmpty();\n"
             "    } catch (Exception e) {\n"
             "        startupError = e;\n"
             "    }\n"
             "}"),
    ("p", "What it means: JavaFX calls init() before it shows anything. The tables are created "
          "if they are missing, the service is built, and example data is added only when there "
          "is no board yet. A problem is not thrown away: it is remembered in startupError and "
          "shown in a dialog after the window opens, so the user sees a reason instead of a "
          "program that closes silently."),
    ("h3", "2. Building the window (MainView.java)"),
    ("code", "public MainView(BoardService service, String databaseLocation) {\n"
             "    this.service = service;\n"
             "    this.databaseLocation = databaseLocation;\n"
             "\n"
             "    getStyleClass().add(\"app-root\");\n"
             "    setTop(buildHeaderArea());\n"
             "    setCenter(buildBoardArea());\n"
             "    setBottom(buildStatusBar());\n"
             "    wireActions();\n"
             "}"),
    ("p", "What it means: MainView is a BorderPane. The top holds the toolbar and the board "
          "title, the centre holds the three columns and the bottom holds the status bar. "
          "wireActions() connects every button and selector to its method, which keeps the "
          "layout code and the behaviour apart."),
    ("h3", "3. One action, end to end (MainView.addTask)"),
    ("code", "Dialogs.showTaskDialog(owner(), draft, currentBoard.getTheme()).ifPresent(edited -> {\n"
             "    try {\n"
             "        service.createTask(currentBoard.getId(), edited.getTitle(), edited.getDescription(),\n"
             "                edited.getPriority(), edited.getStatus(), edited.getDueDate());\n"
             "        refreshBoard();\n"
             "        setStatus(\"Task added\");\n"
             "    } catch (DataAccessException e) {\n"
             "        Dialogs.error(owner(), \"Could not add the task\", e.getMessage());\n"
             "    }\n"
             "});"),
    ("p", "What it means: the dialog returns a result only if the user pressed Save, so "
          "\"cancel\" needs no special case. The service saves the task, the screen reloads, "
          "and a database problem is shown in a readable dialog instead of a stack trace. The "
          "result of the dialog is an Optional, which is why ifPresent() is used."),
    ("h3", "4. Reloading the board (MainView.refreshBoard)"),
    ("code", "for (BoardColumnView columnView : columnViews) {\n"
             "    columnView.setTasks(service.getTasks(currentBoard.getId(), columnView.getStatus()));\n"
             "}"),
    ("p", "What it means: after every change, the three columns are reloaded from the database. "
          "The screen therefore always shows what was really saved, instead of keeping its own "
          "copy that could drift away from the database."),
    ("h3", "5. What a card responds to (TaskCardView.java)"),
    ("code", "setOnMouseClicked(event -> {\n"
             "    if (event.getClickCount() == 2) {\n"
             "        onEdit.accept(task);\n"
             "    }\n"
             "});\n"
             "\n"
             "MenuItem editItem = new MenuItem(\"Edit task\");\n"
             "MenuItem descriptionItem = new MenuItem(\"Show description\");\n"
             "MenuItem deleteItem = new MenuItem(\"Delete task\");\n"
             "ContextMenu contextMenu = new ContextMenu(editItem, descriptionItem, new SeparatorMenuItem(), deleteItem);\n"
             "setOnContextMenuRequested(event -> contextMenu.show(this, event.getScreenX(), event.getScreenY()));"),
    ("p", "What it means: a double click opens the edit form and a right click shows the menu "
          "with the description, editing and deleting. The card only reports what the user did; "
          "the actual work is done by MainView and the service, which keeps the card small."),
    ("h3", "If the examiner asks ..."),
    ("table", [
        ("Question", "Your answer"),
        ("Why JavaFX?",
         "It is the desktop toolkit that comes with the Java platform: it has the controls we need, supports CSS for the themes and has built-in support for dragging and dropping."),
        ("Why is the main class a separate Launcher?",
         "Java refuses to start a class that extends Application from the class path; Launcher is a plain class whose main calls MainApp.main, and the same class is used by the packaged application, the JAR and the tests."),
        ("What happens in the first second?",
         "init() opens the database and creates the service, then start() builds the scene, applies the saved theme and shows the board that was open last time."),
        ("Does the screen block while saving?",
         "No. Everything runs on the JavaFX thread and the queries are small and local, so no background thread is needed; a failure would show an error dialog."),
        ("How is the interface tested?",
         "MainViewTest builds the real window, changes the theme and moves a task; on a machine without a screen the interface tests skip themselves and the other tests still run."),
    ], [3.4, 6.2]),
    ("h3", "Terms to know"),
    ("bullets", [
        "JavaFX - the Java library for desktop windows and controls.",
        "Stage and Scene - the window and the content shown inside it.",
        "Event handler - the code that runs when the user clicks, drags or selects something.",
        "Layout container - a pane such as BorderPane, HBox or VBox that arranges its children.",
        "Optional - a Java container that says \"there may be no result\", used here for the dialogs.",
    ]),
    ("small", "The strongest sentence for the examiner: \"The view only calls the service; SQL "
              "exists in exactly one layer.\""),

    # ------------------------------------------------------------------ M5
    ("h2", "Member 5 - Dragging cards, themes and documentation"),
    ("p", "Open these files: ui/BoardColumnView.java, ui/ThemeManager.java, ui/ThemeListCell.java, "
          "resources/css/base.css and the six theme files, README and docs/."),
    ("say", "Say this first",
     "My part is moving the cards and the appearance. BoardColumnView accepts a dropped card "
     "and works out the position it was dropped at, and ThemeManager swaps the stylesheets so "
     "that six themes are possible. The base CSS file holds all the layout and every colour "
     "comes from a theme file. I also wrote the user and developer documentation and took the "
     "screenshots."),
    ("h3", "1. Starting a drag (TaskCardView.java)"),
    ("code", "setOnDragDetected(event -> {\n"
             "    Dragboard dragboard = startDragAndDrop(TransferMode.MOVE);\n"
             "    ClipboardContent content = new ClipboardContent();\n"
             "    content.put(TASK_DATA_FORMAT, task.getId());\n"
             "    dragboard.setContent(content);\n"
             "    getStyleClass().add(\"task-card-dragging\");\n"
             "    event.consume();\n"
             "});"),
    ("p", "What it means: JavaFX drag and drop moves data through a Dragboard. Only the task "
          "id is carried, not the whole task, so the receiving column can look the task up "
          "again and the data cannot become stale. A CSS class dims the card while it is in "
          "the air."),
    ("h3", "2. Accepting the card (BoardColumnView.java)"),
    ("code", "setOnDragOver(event -> {\n"
             "    if (isTaskDrag(event)) {\n"
             "        event.acceptTransferModes(TransferMode.MOVE);\n"
             "        if (!getStyleClass().contains(\"drop-active\")) {\n"
             "            getStyleClass().add(\"drop-active\");\n"
             "        }\n"
             "    }\n"
             "    event.consume();\n"
             "});\n"
             "\n"
             "setOnDragDropped(event -> {\n"
             "    getStyleClass().remove(\"drop-active\");\n"
             "    long taskId = TaskCardView.draggedTaskId(event.getDragboard());\n"
             "    boolean moved = false;\n"
             "    if (taskId > 0 && dropListener != null) {\n"
             "        dropListener.taskDropped(taskId, status, computeDropIndex(event, taskId));\n"
             "        moved = true;\n"
             "    }\n"
             "    event.setDropCompleted(moved);\n"
             "    event.consume();\n"
             "});"),
    ("p", "What it means: the column first checks whether the dragged data is a task from this "
          "application, and only then accepts a move and highlights itself. On the drop it "
          "reads the id, works out the position, and hands both to the listener - which is "
          "MainView, which calls the service. The column itself saves nothing."),
    ("h3", "3. Where the card lands (BoardColumnView.computeDropIndex)"),
    ("code", "Point2D pointer = cardBox.sceneToLocal(event.getSceneX(), event.getSceneY());\n"
             "List<Node> cards = cardBox.getChildren();\n"
             "\n"
             "int index = 0;\n"
             "for (Node node : cards) {\n"
             "    if (node instanceof TaskCardView card && card.getTask().getId() == draggedTaskId) {\n"
             "        continue;\n"
             "    }\n"
             "    Bounds bounds = node.getBoundsInParent();\n"
             "    if (pointer.getY() < bounds.getMinY() + bounds.getHeight() / 2) {\n"
             "        return index;\n"
             "    }\n"
             "    index++;\n"
             "}\n"
             "return index;"),
    ("p", "What it means: the pointer is compared with the middle of each card. The card is "
          "inserted before the first card whose middle is below the pointer, and at the end "
          "otherwise. The dragged card is skipped, because it is still in the list while it "
          "is being dragged - that is what makes dropping exactly where you aim possible."),
    ("h3", "4. Six themes with eight lines of CSS each (ThemeManager.java)"),
    ("code", "public void applyTheme(Theme theme) {\n"
             "    currentTheme = theme == null ? Theme.MINIMAL_LIGHT : theme;\n"
             "    if (scene != null) {\n"
             "        scene.getStylesheets().setAll(stylesheetsFor(currentTheme));\n"
             "    }\n"
             "}\n"
             "\n"
             "public static List<String> stylesheetsFor(Theme theme) {\n"
             "    List<String> stylesheets = new ArrayList<>(2);\n"
             "    addIfPresent(stylesheets, BASE_STYLESHEET);\n"
             "    addIfPresent(stylesheets, theme.getCssResource());\n"
             "    return stylesheets;\n"
             "}"),
    ("code", "/* dark-mode.css */\n"
             ".root {\n"
             "    -app-bg: #1b1d21;\n"
             "    -app-surface: #202329;\n"
             "    -app-text: #e6e8eb;\n"
             "    -app-accent: #4f9cf9;\n"
             "    -app-priority-high: #f87171;\n"
             "}"),
    ("p", "What it means: base.css holds every layout rule and not one hard-coded colour; it "
          "only uses variables such as -app-card. A theme file defines those variables, so a "
          "new theme is a short file with colours and no layout work at all - which is why six "
          "themes cost almost nothing. The same stylesheets are applied to the dialogs, so a "
          "form looks like the rest of the window. The chosen theme is stored per board."),
    ("h3", "If the examiner asks ..."),
    ("table", [
        ("Question", "Your answer"),
        ("How do you know where the card was dropped?",
         "The pointer is compared with the vertical middle of every card in the column; the dragged card is skipped, so a card can also be dropped where it already is."),
        ("How can six themes be so little code?",
         "The layout file has no colours, only variables, and each theme supplies the values. Adding a theme means adding one small CSS file and one entry in the theme list."),
        ("Why is the theme saved with the board?",
         "Different boards are often different kinds of work; keeping the theme per board means each one can have the look that suits it, and it is restored when the board is opened."),
        ("How did you verify the drag and drop?",
         "MainViewTest builds the real window, performs the move through the same code the mouse uses and checks the new status and order; those tests run in the build with a virtual screen."),
        ("What is in the documentation?",
         "A README for the front page, a guide for users, a build guide for developers, an architecture note, and the screenshots used in the report."),
    ], [3.4, 6.2]),
    ("h3", "Terms to know"),
    ("bullets", [
        "Dragboard - the object that carries the dragged data; here it carries the task id.",
        "Drop target - the area that accepts a dragged item, which is one board column.",
        "CSS stylesheet - the rules for colours, borders and spacing; JavaFX uses a subset of CSS.",
        "CSS variable - a named value such as -app-accent that a theme file defines once.",
        "ThemeManager - the class that applies the base file plus the selected theme to the window.",
    ]),
    ("small", "When the examiner asks about themes, open base.css and show that it contains no "
              "hard-coded colour - that is the whole idea."),

    ("h1", "Part 3 - Questions anybody may be asked"),
    ("table", [
        ("Question", "Short answer", "Add this if the examiner wants more"),
        ("What is CanvasKanban?",
         "A desktop Kanban board for one person, with three columns and cards that can be dragged.",
         "The data is a local SQLite file, so it works without a server or an internet connection."),
        ("Why JavaFX?",
         "It gives desktop controls, CSS theming and drag and drop in the Java platform itself.",
         "The interface is one class tree: MainApp, MainView, the columns, the cards and the dialogs."),
        ("Why SQLite?",
         "The data is one file with no server, which is exactly right for a single-user desktop program.",
         "It supports foreign keys and transactions, which the move operation uses."),
        ("Why Maven?",
         "It records the library versions and gives one command that compiles, tests and packages.",
         "The same command runs in GitHub Actions, so the build is the same everywhere."),
        ("What is the difference between the UI, the service and the DAO?",
         "The UI shows and collects; the service applies the rules; the DAO runs the SQL.",
         "The UI never contains SQL, which is why the rules can be tested without a window."),
        ("Why prepared statements?",
         "The SQL text and the values stay separate, so values cannot break the command.",
         "It also handles types and quotes correctly and reuses the parsed statement."),
        ("Why a transaction when a card is moved?",
         "Two columns change in one action and must be saved together.",
         "If anything fails, the rollback leaves the previous order untouched."),
        ("How are the themes implemented?",
         "One base CSS file for the layout and one small file per theme for the colours.",
         "ThemeManager loads both and the chosen theme is saved with the board."),
        ("What happens if the program is closed?",
         "Nothing is lost: boards, tasks and settings are already in the SQLite file.",
         "On the next start the last opened board and its theme come back."),
        ("Does everybody explain all 22 classes?",
         "No. Each member explains the files of their own area and knows how the layers connect.",
         "Part 1.2 of these notes is the order in which the whole program can be explained."),
    ], [2.9, 3.7, 3.0]),

    ("h1", "Part 4 - Checklist"),
    ("h2", "The night before"),
    ("bullets", [
        "Run \"mvn test\" once and leave the green result on the screen.",
        "Start the application once and make sure the example board and all three columns appear.",
        "Open the six files of your own area and find the lines used in these notes.",
        "Practise your two sentences of \"Say this first\" out loud, without reading them.",
        "Check that the names and enrollment numbers are filled in on the cover of the report.",
    ]),
    ("h2", "On the day"),
    ("bullets", [
        "Have the repository open on the Releases page or the Actions tab as evidence.",
        "Speak from the top down: what the project is, then the layers, then your own module.",
        "Point at the screen while you talk - the sections and tables in the report are ordered for that.",
        "If you do not know an answer, say which layer would contain it and offer to look it up; do not guess.",
        "Only claim the work that is listed for your own member number.",
    ]),
]

NOTES_SUBTITLE = "Viva Speaking Notes for All Five Team Members"

# The picture of the layers, used by both documents and by both formats.
ARCHITECTURE_STEPS = [
    ("MainApp", "starts JavaFX, prepares the database and the service"),
    ("MainView", "toolbar, three columns, cards, dialogs, status bar"),
    ("BoardService", "the application rules - the only thing the screen calls"),
    ("BoardDao / TaskDao / SettingDao", "the SQL, with prepared statements"),
    ("Database", "opens the SQLite connection, switches foreign keys on"),
    ("canvaskanban.db", "one SQLite file in the user's data folder"),
]


# ---------------------------------------------------------------------------
# 5. Mounting the documents (cover pages)
# ---------------------------------------------------------------------------

def cover(subtitle, extra_lines):
    rows = [("Project details", "Information"),
            ("Project title", REPORT_TITLE),
            ("Course", DETAILS["course"]),
            ("College / University", DETAILS["college"]),
            ("Department", DETAILS["department"]),
            ("Guide / Faculty", DETAILS["faculty"]),
            ("Academic year", DETAILS["year"]),
            ("Submission date", DETAILS["date"])]
    blocks = [
        ("cover_title", REPORT_TITLE),
        ("cover_subtitle", subtitle),
        ("cover_space",),
        ("table", rows, [2.6, 6.4]),
        ("cover_space",),
        ("h3", "Submitted by"),
        ("table", [("No.", "Full name", "Enrollment number", "Role")]
                  + [(str(i + 1), n, e, r) for i, (n, e, r) in enumerate(TEAM)],
         [0.5, 2.5, 2.4, 3.6]),
    ]
    for line in extra_lines:
        blocks.append(("small", line))
    blocks.append(("pagebreak",))
    return blocks


def report_blocks():
    return cover(REPORT_SUBTITLE, [
        "Replace every detail printed in dark red before submitting, and keep the same names "
        "and enrollment numbers on the cover of the main report.",
        "Companion document: \"CanvasKanban - Viva Speaking Notes\".",
    ]) + REPORT


def notes_blocks():
    return cover(NOTES_SUBTITLE, [
        "Write your own name into the member line that belongs to your area; the member "
        "numbers are the same role labels that are used in the project report.",
        "Companion document: \"CanvasKanban - Simplified Project Report\".",
    ]) + NOTES


# ---------------------------------------------------------------------------
# 6. The PDF renderer
# ---------------------------------------------------------------------------

FONT = "Helvetica"
FONT_BOLD = "Helvetica-Bold"
MONO = "Courier"

ACCENT = "#12324f"
ACCENT2 = "#1d4e74"
ACCENT3 = "#2d6f9f"
RULE = "#b9c6d2"
HEAD_BG = "#e8eef4"
CODE_BG = "#f4f6f8"
SAY_BG = "#fff6e0"
SAY_BORDER = "#e0b25c"


def _register_fonts():
    """Embeds Noto Sans (and Noto Sans Mono) so the PDF looks the same everywhere."""
    global FONT, FONT_BOLD, MONO
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    sans_dir = "/usr/share/fonts/google-noto"
    files = {n: os.path.join(sans_dir, "NotoSans-%s.ttf" % n)
             for n in ("Regular", "Bold", "Italic", "BoldItalic")}
    if all(os.path.exists(p) for p in files.values()):
        pdfmetrics.registerFont(TTFont("NotoSans", files["Regular"]))
        pdfmetrics.registerFont(TTFont("NotoSans-Bold", files["Bold"]))
        pdfmetrics.registerFont(TTFont("NotoSans-Italic", files["Italic"]))
        pdfmetrics.registerFont(TTFont("NotoSans-BoldItalic", files["BoldItalic"]))
        pdfmetrics.registerFontFamily("NotoSans", normal="NotoSans", bold="NotoSans-Bold",
                                      italic="NotoSans-Italic", boldItalic="NotoSans-BoldItalic")
        FONT, FONT_BOLD = "NotoSans", "NotoSans-Bold"

    mono = os.path.join(sans_dir, "NotoSansMono-Regular.ttf")
    if os.path.exists(mono):
        pdfmetrics.registerFont(TTFont("NotoSansMono", mono))
        MONO = "NotoSansMono"
    return FONT


def build_pdf(path, short_title, title, blocks):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether,
                                    ListFlowable, ListItem, PageBreak, PageTemplate,
                                    Paragraph, Preformatted, Spacer, Table, TableStyle)

    _register_fonts()

    left = right = 1.8 * cm
    usable = A4[0] - left - right

    body = ParagraphStyle("body", fontName=FONT, fontSize=10, leading=14.2, spaceAfter=5)
    small = ParagraphStyle("small", parent=body, fontSize=8.6, leading=11.6,
                           textColor=colors.HexColor("#555555"))
    h1 = ParagraphStyle("h1", fontName=FONT_BOLD, fontSize=15.5, leading=19, spaceBefore=13,
                        spaceAfter=6, textColor=colors.HexColor(ACCENT))
    h2 = ParagraphStyle("h2", fontName=FONT_BOLD, fontSize=12.4, leading=16, spaceBefore=11,
                        spaceAfter=5, textColor=colors.HexColor(ACCENT2))
    h3 = ParagraphStyle("h3", fontName=FONT_BOLD, fontSize=10.8, leading=14.5, spaceBefore=9,
                        spaceAfter=4, textColor=colors.HexColor(ACCENT3))
    cell = ParagraphStyle("cell", parent=body, fontSize=8.8, leading=12.2, spaceAfter=0)
    cellb = ParagraphStyle("cellb", parent=cell, fontName=FONT_BOLD)
    say_title = ParagraphStyle("sayt", parent=body, fontName=FONT_BOLD, fontSize=9.6,
                               leading=13, spaceAfter=3, textColor=colors.HexColor("#7a4a00"))
    cover_title = ParagraphStyle("ct", fontName=FONT_BOLD, fontSize=17.5, leading=22,
                                 alignment=TA_CENTER, textColor=colors.HexColor(ACCENT))
    cover_sub = ParagraphStyle("cs", parent=body, fontSize=13, leading=18, alignment=TA_CENTER,
                               textColor=colors.HexColor("#444444"))
    caption = ParagraphStyle("cap", parent=small, fontSize=8.4, leading=11, alignment=TA_CENTER)

    def P(text, style=body):
        return Paragraph(_pdf_text(text), style)

    def bullets(items, ordered=False):
        return ListFlowable(
            [ListItem(P(i)) for i in items],
            bulletType="1" if ordered else "bullet",
            bulletFontName=FONT, bulletFontSize=8, bulletOffsetY=-1,
            leftIndent=15, spaceAfter=6)

    def table(rows, weights, bold_first=False):
        total = float(sum(weights))
        widths = [usable * w / total for w in weights]
        data = []
        for r, row in enumerate(rows):
            data.append([Paragraph(_pdf_text(str(v)),
                                   cellb if (r == 0 or (bold_first and c == 0)) else cell)
                         for c, v in enumerate(row)])
        t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(HEAD_BG)),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(RULE)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    def code_box(text):
        lines = text.split("\n")
        longest = max(len(line) for line in lines) or 1
        # Noto Sans Mono is 0.6 em wide per character; 0.97 leaves a little slack
        # so a long line is never clipped at the edge of the box.
        size = min(8.2, (usable - 14) * 0.97 / (longest * 0.6))
        size = max(6.0, size)
        style = ParagraphStyle("code", fontName=MONO, fontSize=size, leading=size * 1.32,
                               textColor=colors.HexColor("#1c1c1c"))
        # Preformatted draws the text literally (it does not parse XML), so the
        # code must be handed over exactly as it is - escaping it would show
        # "&amp;&amp;" instead of "&&" and "&lt;Long&gt;" instead of "<Long>".
        inner = Preformatted(text, style)
        box = Table([[inner]], colWidths=[usable])
        box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(CODE_BG)),
            ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor(RULE)),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return box

    def say_box(title, text):
        inner = [Paragraph(_pdf_text(title), say_title), P(text)]
        box = Table([[inner]], colWidths=[usable])
        box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(SAY_BG)),
            ("LINEBEFORE", (0, 0), (0, -1), 2.4, colors.HexColor(SAY_BORDER)),
            ("BOX", (0, 0), (-1, -1), 0.3, colors.HexColor(SAY_BORDER)),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        return box

    def diagram():
        steps = ARCHITECTURE_STEPS
        box_cell = ParagraphStyle("boxcell", parent=body, fontSize=8.8, leading=12.4,
                                  textColor=colors.HexColor("#444444"))
        # The arrow only exists in Noto Sans Mono on this system, so the one glyph
        # is drawn with that font; with the built-in Courier it falls back to "v".
        arrow_cell = ParagraphStyle("arrowcell", parent=body,
                                    fontName=MONO if MONO != "Courier" else FONT,
                                    fontSize=9, leading=12, alignment=TA_CENTER,
                                    textColor=colors.HexColor("#8fa6b8"))
        arrow_glyph = "&#8595;" if MONO != "Courier" else "v"
        rows, commands = [], []
        for i, (name, note) in enumerate(steps):
            if i:                                  # a plain row holding the arrow
                commands.append(("ROWHEIGHT", (0, len(rows)), (-1, len(rows)), 13))
                rows.append([Paragraph(arrow_glyph, arrow_cell)])
            row = len(rows)
            rows.append([Paragraph("<b>%s</b><br/>%s" % (_escape(name), _escape(note)),
                                   box_cell)])
            commands.append(("BACKGROUND", (0, row), (-1, row),
                             colors.HexColor("#eef4fa" if i % 2 == 0 else "#f7fafc")))
            commands.append(("BOX", (0, row), (-1, row), 0.5, colors.HexColor(RULE)))
        picture = Table(rows, colWidths=[usable * 0.92], hAlign="CENTER")
        picture.setStyle(TableStyle(commands + [
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return picture

    def image_grid(items):
        rows, row = [], []
        for entry in items:
            row.append(entry)
            if len(row) == 2:
                rows.append(row)
                row = []
        if row:
            rows.append(row + [None])
        cells = []
        for pair in rows:
            cells.append([image_cell(pair[0], caption),
                          image_cell(pair[1], caption) if pair[1] else ""])
        t = Table(cells, colWidths=[usable / 2, usable / 2])
        t.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        return t

    def image_cell(entry, caption_style):
        name, text = entry
        source = SCREENSHOTS / name
        if not source.exists():
            return P("(missing screenshot: %s)" % name, small)
        from reportlab.lib.utils import ImageReader
        iw, ih = ImageReader(str(source)).getSize()
        width = display_width(source, usable / 2 - 12)
        height = width * ih / iw
        picture = Image(str(source), width=width, height=height)
        picture.hAlign = "CENTER"
        return [picture, Spacer(1, 3), Paragraph(_pdf_text(text), caption_style)]

    def make_footer(canvas, doc):
        if doc.page == 1:
            return
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor(RULE))
        canvas.setLineWidth(0.4)
        canvas.line(left, 1.28 * cm, A4[0] - right, 1.28 * cm)
        canvas.setFont(FONT, 7.6)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(left, 0.92 * cm, short_title)
        canvas.drawRightString(A4[0] - right, 0.92 * cm, "Page %d" % (doc.page - 1))
        canvas.restoreState()

    story = []
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            story.append(P(block[1], h1))
            story.append(_rule(colors, Table, TableStyle, usable))
        elif kind == "h2":
            story.append(P(block[1], h2))
        elif kind == "h3":
            story.append(P(block[1], h3))
        elif kind == "p":
            story.append(P(block[1]))
        elif kind == "small":
            story.append(P(block[1], small))
        elif kind == "bullets":
            story.append(bullets(block[1]))
        elif kind == "steps":
            story.append(bullets(block[1], ordered=True))
        elif kind == "table":
            story.append(table(block[1], block[2]))
            story.append(Spacer(1, 6))
        elif kind == "tableb":
            story.append(table(block[1], block[2], bold_first=True))
            story.append(Spacer(1, 6))
        elif kind == "code":
            story.append(code_box(block[1]))
            story.append(Spacer(1, 4))
        elif kind == "say":
            story.append(say_box(block[1], block[2]))
            story.append(Spacer(1, 8))
        elif kind == "diagram":
            story.append(Spacer(1, 4))
            story.append(diagram())
            story.append(Spacer(1, 10))
        elif kind == "images":
            story.append(image_grid(block[1]))
            story.append(Spacer(1, 4))
        elif kind == "pagebreak":
            story.append(PageBreak())
        elif kind == "cover_title":
            story.append(Spacer(1, 3.4 * cm))
            story.append(P(block[1], cover_title))
        elif kind == "cover_subtitle":
            story.append(Paragraph(_pdf_text(block[1]), cover_sub))
        elif kind == "cover_space":
            story.append(Spacer(1, 1.3 * cm))
        else:
            raise ValueError("unknown block %r" % (kind,))

    doc = BaseDocTemplate(str(path), pagesize=A4, leftMargin=left, rightMargin=right,
                          topMargin=1.6 * cm, bottomMargin=1.7 * cm,
                          title=title, author="CanvasKanban team", subject=short_title)
    frame = Frame(left, 1.7 * cm, usable, A4[1] - 1.6 * cm - 1.7 * cm, id="body")
    doc.addPageTemplates([PageTemplate(id="page", frames=[frame], onPage=make_footer)])
    doc.build(story)
    return doc.page


def _rule(colors, Table, TableStyle, width):
    from reportlab.platypus import Table as _T, TableStyle as _S
    line = _T([[""]], colWidths=[width], rowHeights=[0.1])
    line.setStyle(_S([("LINEABOVE", (0, 0), (-1, 0), 0.9, colors.HexColor(ACCENT))]))
    return line


# ---------------------------------------------------------------------------
# 7. The Word renderer
# ---------------------------------------------------------------------------

def build_docx(path, title, blocks):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    for name, size, colour in (("Heading 1", 15, ACCENT), ("Heading 2", 12.5, ACCENT2),
                               ("Heading 3", 11, ACCENT3)):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(colour.lstrip("#").upper())

    def shade(element, fill):
        props = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") \
            else element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        props.append(shd)

    def add_text(paragraph, text, bold=False, mono=False, size=None):
        """Writes the text, highlighting [placeholders] so they are easy to find."""
        position = 0
        for match in PLACEHOLDER.finditer(text):
            if match.start() > position:
                run = paragraph.add_run(text[position:match.start()])
                run.bold, run.font.name = bold, "Consolas" if mono else run.font.name
                if size:
                    run.font.size = Pt(size)
            run = paragraph.add_run(match.group(0))
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("B00020")
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if mono:
                run.font.name = "Consolas"
            if size:
                run.font.size = Pt(size)
            position = match.end()
        if position < len(text):
            run = paragraph.add_run(text[position:])
            run.bold = bold
            if mono:
                run.font.name = "Consolas"
            if size:
                run.font.size = Pt(size)
        return paragraph

    def paragraph(text, style=None, size=None, colour=None, align=None, space_after=6):
        p = document.add_paragraph(style=style)
        add_text(p, text, size=size)
        if colour:
            for run in p.runs:
                if not run.font.color.rgb:
                    run.font.color.rgb = RGBColor.from_string(colour.lstrip("#").upper())
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def table(rows, weights, bold_first=False):
        t = document.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        total = float(sum(weights))
        available = 17.0
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = t.cell(r, c)
                cell.width = Cm(available * weights[c] / total)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                add_text(p, str(value), bold=(r == 0 or (bold_first and c == 0)), size=9.5)
                if r == 0:
                    shade(cell._tc, "E8EEF4")
        return t

    def code_box(text):
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Cm(0.2)
        lines = text.split("\n")
        for i, line in enumerate(lines):
            if i:
                p.add_run().add_break(WD_BREAK.LINE)
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        shade(p._p, "F4F6F8")
        return p

    def say_box(title, text):
        t = document.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cell = t.cell(0, 0)
        first = cell.paragraphs[0]
        run = first.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("7A4A00")
        body = cell.add_paragraph()
        add_text(body, text)
        shade(cell._tc, "FFF6E0")
        document.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def diagram():
        # The same boxed picture as in the PDF, built from small one-cell tables.
        for i, (name, note) in enumerate(ARCHITECTURE_STEPS):
            if i:
                arrow = document.add_paragraph()
                arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
                arrow.paragraph_format.space_before = Pt(0)
                arrow.paragraph_format.space_after = Pt(0)
                run = arrow.add_run("\u2193")
                run.font.color.rgb = RGBColor.from_string("8FA6B8")
            t = document.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            cell = t.cell(0, 0)
            first = cell.paragraphs[0]
            first.paragraph_format.space_after = Pt(0)
            run = first.add_run(name)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(ACCENT.lstrip("#").upper())
            second = cell.add_paragraph()
            second.paragraph_format.space_after = Pt(0)
            run = second.add_run(note)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string("444444")
            shade(cell._tc, "EEF4FA" if i % 2 == 0 else "F7FAFC")
        document.add_paragraph().paragraph_format.space_after = Pt(4)

    def image_grid(items):
        rows = [items[i:i + 2] for i in range(0, len(items), 2)]
        t = document.add_table(rows=len(rows), cols=2)
        for r, pair in enumerate(rows):
            for c in range(2):
                cell = t.cell(r, c)
                entry = pair[c] if c < len(pair) else None
                p = cell.paragraphs[0]
                if entry is None:
                    continue
                name, text = entry
                source = SCREENSHOTS / name
                if not source.exists():
                    add_text(p, "(missing screenshot: %s)" % name, size=9)
                    continue
                run = p.add_run()
                run.add_picture(str(source),
                                width=Cm(display_width(source, 7.6 * 72 / 2.54) / 72 * 2.54))
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = cell.add_paragraph()
                add_text(cap, text, size=8.5)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return t

    for block in blocks:
        kind = block[0]
        if kind == "h1":
            document.add_heading(block[1], level=1)
        elif kind == "h2":
            document.add_heading(block[1], level=2)
        elif kind == "h3":
            document.add_heading(block[1], level=3)
        elif kind == "p":
            paragraph(block[1])
        elif kind == "small":
            p = paragraph(block[1], size=9)
            for run in p.runs:
                run.font.color.rgb = RGBColor.from_string("555555")
        elif kind == "bullets":
            for item in block[1]:
                paragraph(item, style="List Bullet", space_after=2)
        elif kind == "steps":
            # Written out by hand: Word's "List Number" style would continue the
            # count across every numbered list in the document.
            for i, item in enumerate(block[1], start=1):
                p = paragraph("%d.  %s" % (i, item), space_after=2)
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.75)
        elif kind == "table":
            table(block[1], block[2])
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "tableb":
            table(block[1], block[2], bold_first=True)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "code":
            code_box(block[1])
        elif kind == "say":
            say_box(block[1], block[2])
        elif kind == "diagram":
            diagram()
        elif kind == "images":
            image_grid(block[1])
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "pagebreak":
            document.add_page_break()
        elif kind == "cover_title":
            p = paragraph(block[1], size=20, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(ACCENT.lstrip("#").upper())
        elif kind == "cover_subtitle":
            p = paragraph(block[1], size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
            for run in p.runs:
                run.font.color.rgb = RGBColor.from_string("444444")
        elif kind == "cover_space":
            document.add_paragraph()
        else:
            raise ValueError("unknown block %r" % (kind,))

    document.core_properties.title = title
    document.core_properties.author = "CanvasKanban team"
    document.save(str(path))


# ---------------------------------------------------------------------------
# 8. Build everything
# ---------------------------------------------------------------------------

DOCUMENTS = [
    ("CanvasKanban-Simplified-Report", "CanvasKanban | Simplified project report",
     REPORT_TITLE + " - " + REPORT_SUBTITLE, report_blocks),
    ("CanvasKanban-Viva-Speaking-Notes", "CanvasKanban | Viva speaking notes",
     "CanvasKanban - " + NOTES_SUBTITLE, notes_blocks),
]

# The names the two documents had when they were handed in, so a copy can
# replace them without any confusion.
DOWNLOAD_NAMES = {
    "CanvasKanban-Simplified-Report": "CanvasKanban_Simplified_Report",
    "CanvasKanban-Viva-Speaking-Notes": "CanvasKanban_Code_Explanation_All_Members",
}


def main():
    here = HERE
    download_dir = Path(os.path.expanduser("~/Downloads"))
    for base, short_title, title, content in DOCUMENTS:
        blocks = content()

        placeholders = sum(len(PLACEHOLDER.findall(_plain(block)))
                           for block in blocks if _plain(block))
        code_blocks = [b for b in blocks if b[0] == "code"]

        pdf = here / (base + ".pdf")
        docx = here / (base + ".docx")
        pages = build_pdf(pdf, short_title, title, blocks)
        build_docx(docx, title, blocks)

        print("%-34s %2d pages   %2d code boxes   %2d placeholders"
              % (base + ".pdf", pages, len(code_blocks), placeholders))
        missing = check_glyphs(blocks)
        if missing:
            print("%-34s WARNING: no glyph for %s" % ("", ", ".join(missing)))
        if download_dir.is_dir():
            shutil.copy2(pdf, download_dir / (DOWNLOAD_NAMES[base] + ".pdf"))
            shutil.copy2(docx, download_dir / (DOWNLOAD_NAMES[base] + ".docx"))
            print("%-34s copied to %s" % ("", download_dir))

    longest = max((len(line) for block in DOCUMENTS for line in
                   "\n".join(b[1] for b in block[3]() if b[0] == "code").split("\n")),
                  default=0)
    print("longest line inside a code box: %d characters" % longest)


def _plain(block):
    """The flat text of a block, so placeholders can be counted."""
    if block[0] in ("h1", "h2", "h3", "p", "small"):
        return block[1]
    if block[0] in ("bullets", "steps"):
        return " ".join(block[1])
    if block[0] in ("table", "tableb"):
        return " ".join(str(cell) for row in block[1] for cell in row)
    if block[0] in ("say",):
        return block[2]
    return ""


def _cmap_characters(path):
    """The characters a TrueType font can draw, read from its cmap tables."""


    data = Path(path).read_bytes()
    tables = struct.unpack(">H", data[4:6])[0]
    offset = None
    for i in range(tables):
        record = 12 + 16 * i
        if data[record:record + 4] == b"cmap":
            offset = struct.unpack(">I", data[record + 8:record + 12])[0]
            break
    if offset is None:
        return set()

    characters = set()
    subtables = struct.unpack(">H", data[offset + 2:offset + 4])[0]
    for i in range(subtables):
        record = offset + 4 + 8 * i
        sub = offset + struct.unpack(">I", data[record + 4:record + 8])[0]
        fmt = struct.unpack(">H", data[sub:sub + 2])[0]
        if fmt == 4:
            segments = struct.unpack(">H", data[sub + 6:sub + 8])[0] // 2
            ends = struct.unpack(">%dH" % segments, data[sub + 14:sub + 14 + 2 * segments])
            starts_at = sub + 16 + 2 * segments
            starts = struct.unpack(">%dH" % segments, data[starts_at:starts_at + 2 * segments])
            for start, end in zip(starts, ends):
                if start != 0xFFFF:
                    characters.update(range(start, min(end, 0xFFFF) + 1))
        elif fmt == 12:
            groups = struct.unpack(">I", data[sub + 12:sub + 16])[0]
            for g in range(groups):
                start, end = struct.unpack(">II", data[sub + 16 + 12 * g:sub + 24 + 12 * g])
                characters.update(range(start, min(end, 0x10FFFF) + 1))
    return characters


def check_glyphs(blocks):
    """Lists characters that the embedded fonts cannot draw, which would be blank."""
    available = set()
    for name in ("NotoSans-Regular.ttf", "NotoSans-Bold.ttf", "NotoSansMono-Regular.ttf"):
        path = os.path.join("/usr/share/fonts/google-noto", name)
        if os.path.exists(path):
            available |= _cmap_characters(path)
    if not available:
        return []

    used = set()
    for block in blocks:
        used |= set(_plain(block))
        if block[0] == "code":
            used |= set(block[1])
    return ["%s (U+%04X)" % (c, ord(c)) for c in sorted(used)
            if ord(c) > 126 and c not in available]


if __name__ == "__main__":
    main()
